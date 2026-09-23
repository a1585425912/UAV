"""问题一：复用统一航段 CSV，完整枚举逐箱子集并作字典序精确 DP。

运行：python 问题一_统一航段精确组批.py [--smoke S001]
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
from collections import Counter, defaultdict
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
OUT = ROOT / "results" / "问题一_统一航段"
G0 = 9.81
EPS = 1e-10


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def write_csv(path, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def load_inputs():
    drone_rows = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    drones = {}
    for r in drone_rows[2:5]:
        assert r[0] in ("A", "B", "C")
        drones[r[0]] = dict(zip(("mass0", "payload", "volume", "speed", "range0", "range_full", "energy", "reserve", "prep", "load", "handoff", "extra", "climb_speed", "descent_speed", "eta"), map(float, (*r[2:10], *r[10:17]))))
    assert set(drones) == {"A", "B", "C"}
    wb = load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True)
    raw = list(wb.worksheets[1].values)
    assert raw[0][:5] == ("货箱编号", "服务区编号", "物资类型", "单箱质量（kg）", "单箱体积（m³）")
    assert len(raw) == 81 and raw[-1][0] == "S015-FOD-01"
    boxes = defaultdict(list)
    for r in raw[1:]:
        assert r[0] and r[1] and float(r[3]) > 0 and float(r[4]) > 0
        boxes[r[1]].append({"id": r[0], "mass": float(r[3]), "volume": float(r[4]), "type": r[2]})
    assert len({b["id"] for group in boxes.values() for b in group}) == 80
    assert len(boxes) == 15
    routes = {r["服务区"]: r for r in read_csv(ROOT / "results" / "单服务区往返几何参数.csv")}
    assert set(routes) == set(boxes)
    for area, r in routes.items():
        assert math.isclose(float(r["去程水平距离（m）"]), float(r["返程水平距离（m）"]), abs_tol=1e-7), area
        assert math.isclose(float(r["去程巡航海拔（m）"]), float(r["返程巡航海拔（m）"]), abs_tol=1e-7), area
        assert all(float(r[k]) >= 0 for k in ("去程爬升（m）", "返程爬升（m）", "去程下降（m）", "返程下降（m）"))
    return drones, boxes, routes


def trip(drone, route, q):
    Q = drone["payload"]
    assert 0 <= q <= Q + EPS
    d = float(route["去程水平距离（m）"])
    hout = float(route["去程爬升（m）"])
    hback = float(route["返程爬升（m）"])
    lq = drone["range0"] - (drone["range0"] - drone["range_full"]) * (q / Q) ** 1.5
    parts = {
        "outbound_horizontal_kwh": drone["energy"] * d / lq,
        "outbound_climb_kwh": (drone["mass0"] + q) * G0 * hout / (3.6e6 * drone["eta"]),
        "return_horizontal_kwh": drone["energy"] * d / drone["range0"],
        "return_climb_kwh": drone["mass0"] * G0 * hback / (3.6e6 * drone["eta"]),
    }
    parts["total_kwh"] = sum(parts.values())
    parts["return_soc"] = 1 - parts["total_kwh"] / drone["energy"]
    return parts


def operation_time(drone, route, count):
    d = float(route["去程水平距离（m）"])
    climb = float(route["去程爬升（m）"]) + float(route["返程爬升（m）"])
    descend = float(route["去程下降（m）"]) + float(route["返程下降（m）"])
    fly = climb / drone["climb_speed"] + 2 * d / drone["speed"] + descend / drone["descent_speed"]
    return fly + drone["prep"] + count * drone["load"] + drone["handoff"] + count * drone["extra"]


def safe_payload(drone, route):
    limit = (1 - drone["reserve"] / 100) * drone["energy"]
    empty = trip(drone, route, 0)["total_kwh"]
    full = trip(drone, route, drone["payload"])["total_kwh"]
    if empty > limit + EPS:
        return None, "infeasible_even_empty", empty, full
    if full <= limit:
        return drone["payload"], "structural_payload_limit", empty, full
    lo, hi = 0.0, drone["payload"]
    for _ in range(55):
        mid = (lo + hi) / 2
        if trip(drone, route, mid)["total_kwh"] <= limit:
            lo = mid
        else:
            hi = mid
    return lo, "energy_limited", empty, full


def solve_area(area, boxes, drones, route):
    n = len(boxes)
    total_masks = 1 << n
    weights, volumes = [0.0] * total_masks, [0.0] * total_masks
    candidates, best = [], [None] * total_masks
    for mask in range(1, total_masks):
        bit = mask & -mask
        j = bit.bit_length() - 1
        prev = mask ^ bit
        weights[mask] = weights[prev] + boxes[j]["mass"]
        volumes[mask] = volumes[prev] + boxes[j]["volume"]
        ids = ",".join(boxes[k]["id"] for k in range(n) if mask & (1 << k))
        for name, drone in drones.items():
            if weights[mask] > drone["payload"] + EPS or volumes[mask] > drone["volume"] + EPS:
                continue
            energy = trip(drone, route, weights[mask])
            if energy["total_kwh"] > (1 - drone["reserve"] / 100) * drone["energy"] + EPS:
                continue
            t = operation_time(drone, route, mask.bit_count())
            row = {"服务区": area, "子集掩码": mask, "货箱编号列表": ids, "机型": name,
                   "箱数": mask.bit_count(), "质量_kg": weights[mask], "体积_m3": volumes[mask],
                   "能耗_kwh": energy["total_kwh"], "作业时间_s": t, "返航SOC": energy["return_soc"]}
            candidates.append(row)
            score = (energy["total_kwh"], t, name)
            if best[mask] is None or score < best[mask][0]:
                best[mask] = (score, row)

    @lru_cache(None)
    def dp(remaining):
        if remaining == 0:
            return (0, 0.0, 0.0), ()
        anchor = remaining & -remaining
        sub = remaining
        chosen = None
        while sub:
            if sub & anchor and best[sub] is not None:
                tail = dp(remaining ^ sub)
                if tail is not None:
                    row = best[sub][1]
                    score = (tail[0][0] + 1, tail[0][1] + row["能耗_kwh"], tail[0][2] + row["作业时间_s"])
                    if chosen is None or score < chosen[0]:
                        chosen = score, (sub,) + tail[1]
            sub = (sub - 1) & remaining
        return chosen

    optimum = dp(total_masks - 1)
    if optimum is None:
        raise RuntimeError(f"{area}: 存在无法运输的货箱")
    selected = [best[mask][1] for mask in optimum[1]]
    assert len(candidates) == sum(1 for mask in range(1, total_masks) for drone in drones if weights[mask] <= drones[drone]["payload"] + EPS and volumes[mask] <= drones[drone]["volume"] + EPS and trip(drones[drone], route, weights[mask])["total_kwh"] <= (1 - drones[drone]["reserve"] / 100) * drones[drone]["energy"] + EPS)
    assert sum(r["箱数"] for r in selected) == n
    return candidates, selected, dp.cache_info().currsize


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--smoke", choices=[f"S{i:03d}" for i in range(1, 16)])
    args = parser.parse_args()
    drones, boxes, routes = load_inputs()
    areas = [args.smoke] if args.smoke else sorted(boxes)
    caps, details, candidates, selected, states = [], [], [], [], {}
    for area in areas:
        route = routes[area]
        for name, drone in drones.items():
            cap, status, empty, full = safe_payload(drone, route)
            parts = trip(drone, route, cap) if cap is not None else None
            caps.append({"服务区": area, "机型": name, "最大安全载荷_kg": cap, "状态": status,
                         "额定载荷_kg": drone["payload"], "能量上限_kwh": (1 - drone["reserve"] / 100) * drone["energy"]})
            details.append({"服务区": area, "机型": name, "最大安全载荷_kg": cap, "状态": status,
                            "空载能耗_kwh": empty, "满载能耗_kwh": full, **(parts or {})})
        c, s, state_count = solve_area(area, boxes[area], drones, route)
        candidates.extend(c)
        selected.extend(s)
        states[area] = {"货箱数": len(boxes[area]), "候选数": len(c), "DP状态数": state_count, "架次数": len(s)}
    if args.smoke:
        print(json.dumps({"area": args.smoke, "safe_payload": caps, "trips": selected,
                          "candidate_count": len(candidates), "dp_states": states[args.smoke]["DP状态数"]}, ensure_ascii=False, indent=2))
        return
    OUT.mkdir(parents=True, exist_ok=True)
    for name, rows in (("最大安全载荷_45组.csv", caps), ("安全载荷能耗明细.csv", details),
                       ("全部可行组批.csv", candidates), ("最优组批_逐架次.csv", selected)):
        write_csv(OUT / name, rows)
    assigned = [bid for row in selected for bid in row["货箱编号列表"].split(",")]
    expected = [b["id"] for group in boxes.values() for b in group]
    assert Counter(assigned) == Counter(expected) and len(assigned) == 80
    summary = {"geometry_source": "results/单服务区往返几何参数.csv", "geometry_convention": "DEM节点海拔+局部WGS84水平距离",
               "gravity_m_s2": G0, "reserve": 0.2, "solver": "complete subset enumeration + exact anchored bitmask DP",
               "optimality": "proved for independent service areas and lexicographic trips-energy-time",
               "boxes": len(assigned), "trips": len(selected), "energy_kwh": sum(r["能耗_kwh"] for r in selected),
               "operation_s": sum(r["作业时间_s"] for r in selected),
               "minimum_return_soc": min(r["返航SOC"] for r in selected),
               "model_trips": dict(Counter(r["机型"] for r in selected)), "candidate_rows": len(candidates), "areas": states}
    (OUT / "汇总.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
