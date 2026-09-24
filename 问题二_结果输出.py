"""把经过解码器核验的问题二方案写入 CSV、JSON 与官方结果模板。"""
from __future__ import annotations

import csv
import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "问题二_参考口径"


def write_csv(path, rows):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
        writer.writeheader()
        writer.writerows(rows)


def tabulate(plan, data):
    sorties, delivery = [], []
    for row in plan["sorties"]:
        route = "-".join(stop["area"] for stop in row["stops"])
        sorties.append({"架次编号": row["id"], "无人机编号": row["uav"],
                        "机型编号": row["model"], "电池编号": row["battery"],
                        "开始时刻_s": row["start_s"], "访问服务区顺序": route,
                        "返回O01时刻_s": row["return_s"], "架次能耗_kWh": row["energy_kwh"],
                        "返航SOC": row["return_soc"], "电池充满时刻_s": row["charge_end_s"],
                        "总质量_kg": row["mass_kg"], "总体积_m3": row["volume_m3"],
                        "货箱编号列表": ",".join(row["box_ids"])})
        for stop in row["stops"]:
            for bid in stop["ids"]:
                box = data["boxes"][bid]
                delivery.append({"货箱编号": bid, "架次编号": row["id"],
                                 "服务区编号": box["area"], "交付完成时刻_s": plan["deliveries"][bid],
                                 "物资类型": box["type"], "期望送达时间_s": box["due"],
                                 "硬截止时间_s": box["hard"] if box["hard"] is not None else "",
                                 "优先系数": box["weight"],
                                 "加权延误": box["weight"] * max(0, plan["deliveries"][bid] - box["due"])})
    delivery.sort(key=lambda r: r["货箱编号"])
    return sorties, delivery


def fill_template(sorties, delivery, destination):
    source = ROOT / "results" / "问题一_非枚举整数规划" / "结果提交_问题一主方案.xlsx"
    if not source.exists():
        raise FileNotFoundError("请先运行 问题一_主方案模板.py，确保 Q1 工作表采用最新非枚举方案")
    workbook = load_workbook(source)
    sheets = workbook.sheetnames[:]
    for name, data, fields in (("Q2_运输架次", sorties,
                                ("架次编号", "无人机编号", "机型编号", "电池编号", "开始时刻_s",
                                 "访问服务区顺序", "返回O01时刻_s", "架次能耗_kWh")),
                               ("Q2_逐箱交付", delivery,
                                ("货箱编号", "架次编号", "服务区编号", "交付完成时刻_s"))):
        sheet = workbook[name]
        for idx, record in enumerate(data, 2):
            for column, field in enumerate(fields, 1):
                target = sheet.cell(idx, column)
                if idx > 3:
                    sample = sheet.cell(2, column)
                    target._style = copy(sample._style)
                    target.alignment = copy(sample.alignment)
                value = record[field]
                target.value = float(value) if field in {
                    "开始时刻_s", "返回O01时刻_s", "架次能耗_kWh", "交付完成时刻_s"
                } else value
    assert workbook.sheetnames == sheets
    workbook.save(destination)
    verify = load_workbook(destination, read_only=True, data_only=True)
    assert verify.sheetnames == sheets
    assert verify["Q2_运输架次"].max_row == len(sorties) + 1
    assert verify["Q2_逐箱交付"].max_row == len(delivery) + 1
    verify.close()


def save(plan, data, stem, with_template=False):
    OUT.mkdir(parents=True, exist_ok=True)
    sorties, delivery = tabulate(plan, data)
    assert len(delivery) == 80 and len(sorties) == plan["metrics"]["sorties"]
    assert plan["metrics"]["hard_excess_s"] <= 1e-6
    write_csv(OUT / f"{stem}_逐架次.csv", sorties)
    write_csv(OUT / f"{stem}_逐箱交付.csv", delivery)
    (OUT / f"{stem}_完整方案.json").write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")
    if with_template:
        fill_template(sorties, delivery, OUT / f"结果提交_{stem}.xlsx")
    return {"架次数": len(sorties), "箱数": len(delivery), **plan["metrics"]}
