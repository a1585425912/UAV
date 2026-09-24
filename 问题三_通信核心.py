"""原始 DEM、双向链路预算和保守时间区间通信证书。"""
from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from PIL import Image
from openpyxl import load_workbook

from 地理计算 import geodesic_distance
from 问题二_调度核心 import charge_time


ROOT = Path(__file__).resolve().parent
DEM_PATH = ROOT / "数据" / "镇龙乡及周边30米DEM.tif"
COMM_PATH = ROOT / "数据" / "无人机应急物资运输基础数据" / "通信链路参数.xlsx"
NODE_PATH = ROOT / "数据" / "无人机应急物资运输基础数据" / "调度中心与服务区.xlsx"


@dataclass(frozen=True)
class Point:
    lon: float
    lat: float
    alt: float


class Terrain:
    def __init__(self, path=DEM_PATH):
        image = Image.open(path)
        assert image.mode == "F"
        self.z = np.asarray(image, dtype=np.float32)
        self.step_lon, self.step_lat = map(float, image.tag_v2[33550][:2])
        tie = image.tag_v2[33922]
        self.left, self.top = float(tie[3]), float(tie[4])
        assert image.tag_v2[34735][-1] == 4326
        assert self.z.shape == (1309, 1486) and np.isfinite(self.z).all()
        self.cell_x = 111320 * math.cos(math.radians(23.05)) * self.step_lon
        self.cell_y = 111320 * self.step_lat

    def index(self, lon, lat):
        col = math.floor((lon - self.left) / self.step_lon)
        row = math.floor((self.top - lat) / self.step_lat)
        if not (0 <= row < self.z.shape[0] and 0 <= col < self.z.shape[1]):
            raise ValueError(f"端点超出 DEM: {lon}, {lat}")
        return row, col

    def elevation(self, lon, lat):
        return float(self.z[self.index(lon, lat)])

    def path_max(self, start: Point, end: Point):
        a, b = self.xy(start), self.xy(end)
        vec = b - a
        length2 = float(np.dot(vec, vec))
        radius = math.hypot(self.cell_x, self.cell_y) / 2 + .1
        cols = np.arange(max(0, int(math.floor((min(a[0], b[0]) - radius) / self.cell_x))),
                         min(self.z.shape[1], int(math.ceil((max(a[0], b[0]) + radius) / self.cell_x)) + 1))
        rows = np.arange(max(0, int(math.floor((min(a[1], b[1]) - radius) / self.cell_y))),
                         min(self.z.shape[0], int(math.ceil((max(a[1], b[1]) + radius) / self.cell_y)) + 1))
        if not len(rows) or not len(cols):
            raise ValueError("中继航段超出 DEM")
        cc, rr = np.meshgrid(cols, rows)
        x = (cc + .5) * self.cell_x
        y = (rr + .5) * self.cell_y
        u = np.clip(((x - a[0]) * vec[0] + (y - a[1]) * vec[1]) / max(length2, 1e-12), 0, 1)
        distance = np.hypot(x - a[0] - u * vec[0], y - a[1] - u * vec[1])
        mask = distance <= radius
        if not np.any(mask):
            raise ValueError("中继航段未落入 DEM")
        return float(self.z[rr[mask], cc[mask]].max())

    def xy(self, point):
        return np.array([(point.lon - self.left) * 111320 * math.cos(math.radians(23.05)),
                         (self.top - point.lat) * 111320], dtype=float)

    def certify_clear(self, moving0: Point, moving1: Point, fixed: Point):
        """以扫掠线段外包走廊和最低视线高证明整段无 DEM 遮挡。"""
        a0, a1, b = self.xy(moving0), self.xy(moving1), self.xy(fixed)
        mid = (a0 + a1) / 2
        vec = b - mid
        length = float(np.linalg.norm(vec))
        if length < 1e-8:
            return min(moving0.alt, moving1.alt, fixed.alt) > self.elevation(fixed.lon, fixed.lat)
        movement = float(np.linalg.norm(a1 - a0)) / 2
        cell_radius = math.hypot(self.cell_x, self.cell_y) / 2
        radius = movement + cell_radius + 0.5
        xlo, xhi = min(mid[0], b[0]) - radius, max(mid[0], b[0]) + radius
        ylo, yhi = min(mid[1], b[1]) - radius, max(mid[1], b[1]) + radius
        cols = np.arange(max(0, int(math.floor(xlo / self.cell_x))),
                         min(self.z.shape[1], int(math.ceil(xhi / self.cell_x)) + 1))
        rows = np.arange(max(0, int(math.floor(ylo / self.cell_y))),
                         min(self.z.shape[0], int(math.ceil(yhi / self.cell_y)) + 1))
        if not len(cols) or not len(rows):
            return False
        cc, rr = np.meshgrid(cols, rows)
        x = (cc + .5) * self.cell_x
        y = (rr + .5) * self.cell_y
        u = ((x - mid[0]) * vec[0] + (y - mid[1]) * vec[1]) / length**2
        uc = np.clip(u, 0, 1)
        distance = np.hypot(x - (mid[0] + uc * vec[0]), y - (mid[1] + uc * vec[1]))
        mask = distance <= radius
        if not np.any(mask):
            return False
        # 端点运动使投影分数有误差；按全部可能的线段位置保守扩大。
        du = min(1.0, 2 * radius / length)
        ul = np.clip(u - du, 0, 1)
        uh = np.clip(u + du, 0, 1)
        low_start = min(moving0.alt, moving1.alt)
        za = low_start + ul * (fixed.alt - low_start)
        zb = low_start + uh * (fixed.alt - low_start)
        lower = np.minimum(za, zb)
        # 保留端点像元作保守检查，避免近端高地被跳过。
        return bool(np.all(self.z[rr[mask], cc[mask]] < lower[mask] - 0.1))


def load_parameters():
    rows = list(load_workbook(COMM_PATH, read_only=True, data_only=True).active.values)
    values = {str(r[3]): float(r[4]) for r in rows[2:7] if r[3] and r[4] is not None}
    # 发射功率和天线增益按设备端点顺序提取，不混用接入/回传接口。
    devices = {}
    for name, start in (("T", 7), ("RA", 9), ("RB", 11), ("G", 13)):
        devices[name] = (float(rows[start][4]), float(rows[start + 1][4]))
    sensitivity, margin = float(rows[5][4]), float(rows[6][4])
    threshold = sensitivity + margin
    def limit(a, b):
        ab = devices[a][0] + devices[a][1] + devices[b][1] - values["Lsys"] - threshold
        ba = devices[b][0] + devices[b][1] + devices[a][1] - values["Lsys"] - threshold
        return min(ab, ba)
    return {"freq_mhz": values["f"], "obstruction_db": values["Lobs"],
            "limit_direct": limit("T", "G"), "limit_access": limit("T", "RA"),
            "limit_backhaul": limit("RB", "G"), "gateway_height": float(rows[15][4])}


def load_nodes():
    rows = list(load_workbook(NODE_PATH, read_only=True, data_only=True).active.values)
    relevant = [rows[2], *rows[6:21]]
    return {r[0]: Point(float(r[2]), float(r[3]), float(r[4])) for r in relevant}


def maximum_distance(a: Point, b: Point, fixed: Point):
    def length(p):
        horizontal = geodesic_distance(p.lon, p.lat, fixed.lon, fixed.lat)
        return math.hypot(horizontal, p.alt - fixed.alt)
    # 区间内部点到 a/b 的三维位移以全球 WGS84 每度上界保守包住。
    # 这会高估真实距离，但不能把接近门限的链路误判为可用。
    motion_bound = (111320 * abs(b.lon - a.lon) + 111700 * abs(b.lat - a.lat)
                    + abs(b.alt - a.alt))
    return max(length(a), length(b)) + motion_bound + .1


def certified_link(terrain: Terrain, params, moving0: Point, moving1: Point,
                   fixed: Point, threshold_db: float):
    distance_km = max(maximum_distance(moving0, moving1, fixed) / 1000, 1e-9)
    fspl = 32.45 + 20 * math.log10(params["freq_mhz"]) + 20 * math.log10(distance_km)
    if fspl + params["obstruction_db"] <= threshold_db - 1e-9:
        return True, "遮挡上界仍可用"
    if fspl > threshold_db - 1e-9:
        return False, "距离上界未获认证"
    clear = terrain.certify_clear(moving0, moving1, fixed)
    return clear, ("走廊证明无遮挡" if clear else "视线未认证")


def relay_sortie(terrain: Terrain, relay: Point, depart_s: float, service_end_s: float,
                 relay_id: str, energy_id: str):
    rows = list(load_workbook(ROOT / "数据" / "无人机应急物资运输基础数据" / "中继无人机数据.xlsx",
                              read_only=True, data_only=True).active.values)
    raw = rows[2]
    mass, speed, cruise_power, full = map(float, (raw[4], raw[5], raw[6], raw[7]))
    reserve, prep, link, turn = map(float, (raw[8] / 100, raw[9], raw[10], raw[11]))
    climb_speed, descent_speed, eta, hover_power, comm_power = map(
        float, (raw[12], raw[13], raw[14], raw[16], raw[17]))
    hub = load_nodes()["O01"]
    if relay.alt - terrain.elevation(relay.lon, relay.lat) > float(raw[18]) + 1e-8:
        raise ValueError("中继悬停高度超限")
    cruise = max(terrain.path_max(hub, relay) + 50, relay.alt)
    distance = geodesic_distance(hub.lon, hub.lat, relay.lon, relay.lat)
    up_out, down_out = cruise - hub.alt, cruise - relay.alt
    up_back, down_back = cruise - relay.alt, cruise - hub.alt
    out = up_out / climb_speed + distance / speed + down_out / descent_speed
    back = up_back / climb_speed + distance / speed + down_back / descent_speed
    transit_energy = (2 * cruise_power * distance / speed / 3600
                      + mass * 9.81 * (up_out + up_back) / (3.6e6 * eta))
    ready = depart_s + prep + out + link
    if service_end_s < ready:
        raise ValueError("服务结束早于建链")
    ret = service_end_s + back
    # 建链 30 s 期间已在悬停并运行通信模块，故从到位即计双功率。
    service_energy = (hover_power + comm_power) * (service_end_s - (ready - link)) / 3600
    total = transit_energy + service_energy
    soc = 1 - total / full
    if soc < reserve - 1e-8:
        raise ValueError(f"中继返航 SOC 不足: {soc}")
    charge_end = ret + charge_time(soc, float(rows[11][2]))
    return {"relay_id": relay_id, "energy_id": energy_id,
            "depart_s": depart_s, "link_ready_s": ready, "service_end_s": service_end_s,
            "return_s": ret, "relay_free_s": ret + turn, "energy_free_s": charge_end,
            "energy_kwh": total, "soc": soc, "lon": relay.lon, "lat": relay.lat,
            "hover_alt_m": relay.alt, "cruise_alt_m": cruise, "distance_m": distance}


def smoke():
    terrain, params, nodes = Terrain(), load_parameters(), load_nodes()
    gateway = Point(nodes["O01"].lon, nodes["O01"].lat,
                    nodes["O01"].alt + params["gateway_height"])
    near = Point(nodes["O01"].lon, nodes["O01"].lat, nodes["O01"].alt + 30)
    good, reason = certified_link(terrain, params, near, near, gateway, params["limit_direct"])
    assert good
    for area in ("S003", "S004", "S012"):
        place = nodes[area]
        p = Point(place.lon, place.lat, place.alt + 30)
        print(area, certified_link(terrain, params, p, p, gateway, params["limit_direct"]))
    print("链路阈值", params, "同址检查", reason)
    candidates = {"P1": (109.19944, 23.05750), "P2": (109.28028, 23.03333),
                  "P3": (109.19444, 23.06333)}
    for name, (lon, lat) in candidates.items():
        point = Point(lon, lat, terrain.elevation(lon, lat) + 300)
        backhaul = certified_link(terrain, params, point, point, gateway, params["limit_backhaul"])
        sortie = relay_sortie(terrain, point, 0, 2000, "R01", "R-B1")
        print(name, "悬停海拔", point.alt, "回传", backhaul, "建链", sortie["link_ready_s"],
              "能耗", sortie["energy_kwh"])


if __name__ == "__main__":
    smoke()
