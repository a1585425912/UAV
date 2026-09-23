"""从原始货箱/机型工作簿独立复核问题一全部交付结果。"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "问题一_参考口径"
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"


def rows(name):
    with (OUT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def close(actual, expected, label):
    assert math.isclose(float(actual), float(expected), rel_tol=0, abs_tol=1e-7), (label, actual, expected)


def main():
    drone_raw = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    drones = {r[0]: r for r in drone_raw[2:5]}
    box_raw = list(load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True).worksheets[1].values)
    boxes = {r[0]: r for r in box_raw[1:]}
    routes = {r["服务区"]: r for r in rows("单服务区往返几何参数.csv")}
    assert len(drones) == 3 and len(boxes) == 80 and len(routes) == 15

    def energy(model, area, payload):
        r, route = drones[model], routes[area]
        mass0, capacity, range0, range_full, battery, eta = map(float, (r[2], r[3], r[6], r[7], r[8], r[16]))
        loaded_range = range0 - (range0 - range_full) * (payload / capacity) ** 1.5
        return (battery * float(route["去程水平距离（m）"]) / loaded_range
                + battery * float(route["返程水平距离（m）"]) / range0
                + 9.81 * ((mass0 + payload) * float(route["去程爬升（m）"])
                          + mass0 * float(route["返程爬升（m）"])) / (3.6e6 * eta))

    def validate(plan, reserve, label):
        covered = []
        for row in plan:
            area, model = row["服务区"], row["机型"]
            assert area in routes and model in drones
            r, route = drones[model], routes[area]
            ids = row["货箱编号列表"].split(",")
            assert ids and all(bid in boxes and boxes[bid][1] == area for bid in ids), (label, area, ids)
            mass = sum(float(boxes[bid][3]) for bid in ids)
            volume = sum(float(boxes[bid][4]) for bid in ids)
            e = energy(model, area, mass)
            flight = ((float(route["去程爬升（m）"]) + float(route["返程爬升（m）"])) / float(r[14])
                      + (float(route["去程水平距离（m）"]) + float(route["返程水平距离（m）"])) / float(r[5])
                      + (float(route["去程下降（m）"]) + float(route["返程下降（m）"])) / float(r[15]))
            operation = flight + float(r[10]) + len(ids) * float(r[11]) + float(r[12]) + len(ids) * float(r[13])
            roundtrip = operation - float(r[10]) - len(ids) * float(r[11])
            assert mass <= float(r[3]) + 1e-8 and volume <= float(r[4]) + 1e-8
            assert e <= (1 - reserve / 100) * float(r[8]) + 1e-8, (label, area, model, "SOC")
            for key, value in (("箱数", len(ids)), ("质量_kg", mass), ("体积_m3", volume),
                               ("能耗_kwh", e), ("作业时间_s", operation),
                               ("往返时间_含交接_s", roundtrip), ("返航SOC", 1 - e / float(r[8]))):
                close(row[key], value, (label, area, key))
            covered.extend(ids)
        assert Counter(covered) == Counter(boxes.keys()), (label, "货箱未恰好覆盖一次")
        return sum(float(r["能耗_kwh"]) for r in plan), sum(float(r["作业时间_s"]) for r in plan)

    caps = rows("最大安全载荷_45组.csv")
    assert len(caps) == 45 and {(r["服务区"], r["机型"]) for r in caps} == {(a, g) for a in routes for g in drones}
    for row in caps:
        q = float(row["最大安全载荷_kg"])
        capacity = float(drones[row["机型"]][3])
        budget = .8 * float(drones[row["机型"]][8])
        assert 0 <= q <= capacity and energy(row["机型"], row["服务区"], q) <= budget + 1e-8
        if row["状态"] == "energy_limited":
            close(energy(row["机型"], row["服务区"], q), budget, "载荷边界")
        else:
            assert row["状态"] == "structural_payload_limit" and q == capacity

    base = rows("最优组批_逐架次.csv")
    e0, t0 = validate(base, 20, "基准")
    summary = json.loads((OUT / "汇总.json").read_text(encoding="utf-8"))
    assert len(base) == summary["架次数"] == 18
    close(summary["总能耗_kWh"], e0, "基准总能耗")
    close(summary["累计作业时间_s"], t0, "基准时间")
    scenarios = rows("返航余量组批_逐架次.csv")
    sensitivities = rows("返航余量汇总.csv")
    assert len(sensitivities) == 6
    for record in sensitivities:
        reserve = float(record["返航余量_百分比"])
        selected = [row for row in scenarios if float(row["返航余量_百分比"]) == reserve]
        e, t = validate(selected, reserve, f"余量{reserve}")
        assert len(selected) == int(record["架次数"])
        close(e, record["总能耗_kWh"], "余量能耗")
        close(t, record["累计作业时间_s"], "余量时间")
    tradeoffs, tradeoff_rows = rows("架次能耗权衡.csv"), rows("架次能耗权衡_逐架次.csv")
    assert [int(r["总架次数"]) for r in tradeoffs] == list(range(18, 24))
    for record in tradeoffs:
        k = int(record["总架次数"])
        selected = [r for r in tradeoff_rows if int(r["总架次数"]) == k]
        e, t = validate(selected, 20, f"权衡{k}")
        assert len(selected) == k
        close(e, record["最小总能耗_kWh"], "权衡能耗")
        close(t, record["对应累计作业时间_s"], "权衡时间")
    workbook = load_workbook(OUT / "结果提交_问题一参考口径.xlsx", read_only=True, data_only=True)
    assert len(workbook.sheetnames) == 6
    sheet = list(workbook["Q1_单点组批"].values)
    assert len(sheet) == 19
    for idx, (record, row) in enumerate(zip(base, sheet[1:]), 1):
        assert row[:4] == (f"Q1-{idx:02d}", record["服务区"], record["机型"], record["货箱编号列表"])
        for actual, expected in zip(row[4:9], (record["质量_kg"], record["体积_m3"],
                                                  record["往返时间_含交接_s"], record["能耗_kwh"],
                                                  100 * float(record["返航SOC"]))):
            close(actual, expected, "提交模板")
    result = {"状态": "PASS", "货箱": 80, "服务区": 15, "安全载荷": 45,
              "基准架次": 18, "余量场景": 6, "权衡点": 6, "模板行": 18,
              "基准能耗_kWh": e0, "基准时间_s": t0}
    (OUT / "独立审计报告.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=False))


if __name__ == "__main__":
    main()
