"""不调用求解/解码模块，从原始 XLSX 与航段表复算问题二方案。"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "问题二_参考口径"
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"


def read(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def close(actual, expected, label, tol=1e-6):
    assert math.isclose(float(actual), float(expected), abs_tol=tol, rel_tol=0), (label, actual, expected)


def inputs():
    raw = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    models = {r[0]: r for r in raw[2:5]}
    uavs = {r[0]: r[1] for r in raw[8:16]}
    battery = {r[0]: (int(r[1]), float(r[2])) for r in raw[19:22]}
    raw_boxes = list(load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True).worksheets[1].values)
    boxes = {r[0]: r for r in raw_boxes[1:]}
    legs = {(r["起点"], r["终点"]): r for r in read(OUT / "有向航段几何参数.csv")}
    assert len(models) == 3 and len(uavs) == 8 and len(boxes) == 80 and len(legs) == 240
    return models, uavs, battery, boxes, legs


def recharge(soc, full):
    return full * (.65 * (.9 - soc) / .9 + .35) if soc < .9 else full * .35 * (1 - soc) / .1


def leg_value(model, leg, payload):
    mass0, cap, speed, range0, range_full, battery, climb_speed, descent_speed, eta = map(
        float, (model[2], model[3], model[5], model[6], model[7], model[8], model[14], model[15], model[16]))
    distance = float(leg["水平距离（m）"])
    range_loaded = range0 - (range0 - range_full) * (payload / cap) ** 1.5
    energy = battery * distance / range_loaded + (mass0 + payload) * 9.81 * float(leg["起点爬升（m）"]) / (3.6e6 * eta)
    seconds = (float(leg["起点爬升（m）"]) / climb_speed + distance / speed
               + float(leg["终点下降（m）"]) / descent_speed)
    return energy, seconds


def verify_plan(stem, inputs_bundle):
    models, uavs, batteries, boxes, legs = inputs_bundle
    sorties = read(OUT / f"{stem}_逐架次.csv")
    deliveries = read(OUT / f"{stem}_逐箱交付.csv")
    assert len(deliveries) == 80 and len({r["货箱编号"] for r in deliveries}) == 80
    by_trip = defaultdict(list)
    for row in deliveries:
        bid = row["货箱编号"]
        assert bid in boxes and row["服务区编号"] == boxes[bid][1]
        by_trip[row["架次编号"]].append(row)
    assert {r["架次编号"] for r in sorties} == set(by_trip)
    uav_intervals, battery_intervals = defaultdict(list), defaultdict(list)
    hard_excess, weighted, energy_sum, makespan = 0.0, 0.0, 0.0, 0.0
    for row in sorties:
        sid, uid, model_key, battery_id = (row[k] for k in ("架次编号", "无人机编号", "机型编号", "电池编号"))
        model = models[model_key]
        assert uavs[uid] == model_key and battery_id in {f"{model_key}-B{i}" for i in range(1, batteries[model_key][0] + 1)}
        assigned = by_trip[sid]
        ids = [r["货箱编号"] for r in assigned]
        mass = sum(float(boxes[bid][3]) for bid in ids)
        volume = sum(float(boxes[bid][4]) for bid in ids)
        close(mass, row["总质量_kg"], (stem, sid, "质量"))
        close(volume, row["总体积_m3"], (stem, sid, "体积"))
        assert mass <= float(model[3]) + 1e-8 and volume <= float(model[4]) + 1e-8
        areas = row["访问服务区顺序"].split("-")
        assert areas and len(areas) == len(set(areas))
        assert {boxes[bid][1] for bid in ids} == set(areas)
        start = float(row["开始时刻_s"])
        elapsed = float(model[10]) + len(ids) * float(model[11])
        used = 0.0
        remaining, at = mass, "O01"
        for area in areas:
            e, seconds = leg_value(model, legs[at, area], remaining)
            used += e
            elapsed += seconds
            here = [r for r in assigned if r["服务区编号"] == area]
            elapsed += float(model[12]) + len(here) * float(model[13])
            for delivered in here:
                bid = delivered["货箱编号"]
                close(delivered["交付完成时刻_s"], start + elapsed, (stem, sid, bid, "交付时刻"))
                hard_candidates = []
                if boxes[bid][5] == "是" and boxes[bid][6] is not None:
                    hard_candidates.append(float(boxes[bid][6]))
                if boxes[bid][2] == "医疗物资":
                    hard_candidates.append(float(boxes[bid][7]))
                due = float(boxes[bid][7])
                actual = start + elapsed
                if hard_candidates:
                    hard_excess += max(0, actual - min(hard_candidates))
                weighted += float(boxes[bid][8]) * max(0, actual - due)
            remaining -= sum(float(boxes[r["货箱编号"]][3]) for r in here)
            at = area
        assert abs(remaining) < 1e-7
        e, seconds = leg_value(model, legs[at, "O01"], 0.0)
        used += e
        elapsed += seconds
        ret = start + elapsed
        soc = 1 - used / float(model[8])
        assert soc >= float(model[9]) / 100 - 1e-8
        close(used, row["架次能耗_kWh"], (stem, sid, "能耗"))
        close(ret, row["返回O01时刻_s"], (stem, sid, "返回"))
        close(soc, row["返航SOC"], (stem, sid, "SOC"))
        charge_end = ret + recharge(soc, batteries[model_key][1])
        close(charge_end, row["电池充满时刻_s"], (stem, sid, "充电"))
        uav_intervals[uid].append((start, ret, sid))
        battery_intervals[battery_id].append((start, charge_end, sid))
        energy_sum += used
        makespan = max(makespan, ret)
    for kind, resources in (("无人机", uav_intervals), ("电池", battery_intervals)):
        for key, intervals in resources.items():
            for old, new in zip(sorted(intervals), sorted(intervals)[1:]):
                assert old[1] <= new[0] + 1e-8, (stem, kind, key, old, new)
    assert hard_excess <= 1e-6, (stem, "硬时限", hard_excess)
    full = json.loads((OUT / f"{stem}_完整方案.json").read_text(encoding="utf-8"))
    for key, expected in (("hard_excess_s", hard_excess), ("weighted_tardiness", weighted),
                          ("makespan_s", makespan), ("energy_kwh", energy_sum),
                          ("sorties", len(sorties)), ("boxes", 80)):
        close(full["metrics"][key], expected, (stem, key))
    return {"方案": stem, "架次数": len(sorties), "加权延误": weighted,
            "完成时间_s": makespan, "能耗_kWh": energy_sum}


def main():
    source = inputs()
    stems = sorted(path.name.removesuffix("_逐架次.csv") for path in OUT.glob("*_逐架次.csv"))
    assert "主方案" in stems
    records = [verify_plan(stem, source) for stem in stems]
    for stem in ("主方案",):
        path = OUT / f"结果提交_{stem}.xlsx"
        if not path.exists():
            continue
        book = load_workbook(path, read_only=True, data_only=True)
        sorties = read(OUT / f"{stem}_逐架次.csv")
        deliveries = read(OUT / f"{stem}_逐箱交付.csv")
        assert book["Q2_运输架次"].max_row == len(sorties) + 1
        assert book["Q2_逐箱交付"].max_row == len(deliveries) + 1
        assert len(book.sheetnames) == 6
        for s, values in zip(sorties, list(book["Q2_运输架次"].values)[1:]):
            assert values[:4] == tuple(s[k] for k in ("架次编号", "无人机编号", "机型编号", "电池编号"))
            close(values[4], s["开始时刻_s"], (stem, "模板开始时刻"))
            assert values[5] == s["访问服务区顺序"]
            close(values[6], s["返回O01时刻_s"], (stem, "模板返航时刻"))
            close(values[7], s["架次能耗_kWh"], (stem, "模板能耗"))
        for d, values in zip(deliveries, list(book["Q2_逐箱交付"].values)[1:]):
            assert values[:3] == tuple(d[k] for k in ("货箱编号", "架次编号", "服务区编号"))
            close(values[3], d["交付完成时刻_s"], (stem, "模板交付时刻"))
        source_book = load_workbook(ROOT / "results" / "问题一_非枚举整数规划" / "结果提交_问题一主方案.xlsx",
                                    read_only=True, data_only=True)
        for sheet in ("Q1_单点组批", "Q3_中继架次", "Q3_通信保障", "Q4_分区配置"):
            assert list(book[sheet].values) == list(source_book[sheet].values), (stem, sheet, "原工作表变动")
        source_book.close()
        book.close()
    result = {"状态": "PASS", "已审计方案数": len(records), "结果": records}
    (OUT / "独立审计报告.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"状态": "PASS", "已审计方案数": len(records),
                      "主方案": next(r for r in records if r["方案"] == "主方案")}, ensure_ascii=False))


if __name__ == "__main__":
    main()
