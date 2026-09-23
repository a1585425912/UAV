"""只填写官方结果模板的 Q1_单点组批 工作表，其他工作表保持原样。"""
from __future__ import annotations

import csv
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "问题一_参考口径"
DEST = OUT / "结果提交_问题一参考口径.xlsx"
OFFICIAL_TEMPLATE = ROOT.parent / "D题" / "结果提交模板.xlsx"
TEMPLATE = OFFICIAL_TEMPLATE if OFFICIAL_TEMPLATE.exists() else DEST


def main():
    with (OUT / "最优组批_逐架次.csv").open(encoding="utf-8-sig", newline="") as handle:
        trips = list(csv.DictReader(handle))
    assert len(trips) == 18
    workbook = load_workbook(TEMPLATE)
    original_sheets = workbook.sheetnames[:]
    ws = workbook["Q1_单点组批"]
    expected = ("架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）",
                "总体积（m³）", "往返时间（s）", "架次能耗（kWh）", "返航SOC（%）")
    assert tuple(ws.cell(1, column).value for column in range(1, 10)) == expected
    for row_number, trip in enumerate(trips, 2):
        for column in range(1, 10):
            if row_number > 3:
                source, target = ws.cell(2, column), ws.cell(row_number, column)
                if source.has_style:
                    target._style = copy(source._style)
                target.alignment = copy(source.alignment)
        values = (f"Q1-{row_number-1:02d}", trip["服务区"], trip["机型"], trip["货箱编号列表"],
                  float(trip["质量_kg"]), float(trip["体积_m3"]),
                  float(trip["往返时间_含交接_s"]), float(trip["能耗_kwh"]),
                  100 * float(trip["返航SOC"]))
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
    assert workbook.sheetnames == original_sheets
    workbook.save(DEST)
    verify = load_workbook(DEST, read_only=True, data_only=True)
    output_rows = list(verify["Q1_单点组批"].values)
    assert len(output_rows) == 19 and output_rows[1][0] == "Q1-01" and output_rows[-1][0] == "Q1-18"
    assert verify.sheetnames == original_sheets
    print(f"已填写 Q1 18 架次，其他 {len(original_sheets)-1} 张工作表保留：{DEST}")


if __name__ == "__main__":
    main()
