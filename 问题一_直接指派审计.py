"""独立读取原始工作簿与有向航段，复核直接整数规划的逐架次结果。"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "results" / "问题一_非枚举整数规划"
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"


def read_rows(path):
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def main():
    drone_rows = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    drones = {row[0]: row for row in drone_rows[2:5]}
    box_rows = list(load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True).worksheets[1].values)
    boxes = {row[0]: row for row in box_rows[1:]}
    legs = {(row["起点"], row["终点"]): row for row in read_rows(ROOT / "results" / "有向航段几何参数.csv")}
    caps = read_rows(OUT / "最大安全载荷_45组.csv")
    batches = read_rows(OUT / "最优组批_逐架次.csv")
    scenarios = read_rows(OUT / "返航余量组批_逐架次.csv")
    sensitivity = read_rows(OUT / "返航余量敏感性.csv")
    summary = json.loads((OUT / "汇总.json").read_text(encoding="utf-8"))
    assert len(drones) == 3 and len(boxes) == 80 and len(legs) == 240 and len(caps) == 45

    def energy(model, area, mass):
        r = drones[model]
        out, back = legs[("O01", area)], legs[(area, "O01")]
        empty_mass, max_mass, no_load_range, full_range, battery, eta = map(float, (r[2], r[3], r[6], r[7], r[8], r[16]))
        loaded_range = no_load_range - (no_load_range - full_range) * (mass / max_mass) ** 1.5
        return (battery * float(out["水平距离（m）"]) / loaded_range
                + battery * float(back["水平距离（m）"]) / no_load_range
                + 9.81 * ((empty_mass + mass) * float(out["起点爬升（m）"])
                          + empty_mass * float(back["起点爬升（m）"])) / (3.6e6 * eta))

    def check_rows(rows, reserve, label):
        assigned = []
        totals = defaultdict(int)
        for row in rows:
            area, model = row["服务区"], row["机型"]
            r = drones[model]
            ids = row["货箱编号列表"].split(",")
            assert ids and len(ids) == len(set(ids)), (label, area, "重复箱")
            assert all(boxes[bid][1] == area for bid in ids), (label, area, "跨区")
            mass = sum(float(boxes[bid][3]) for bid in ids)
            volume = sum(float(boxes[bid][4]) for bid in ids)
            actual = energy(model, area, mass)
            out, back = legs[("O01", area)], legs[(area, "O01")]
            ascend = float(out["起点爬升（m）"]) + float(back["起点爬升（m）"])
            descend = float(out["终点下降（m）"]) + float(back["终点下降（m）"])
            flight = ascend / float(r[14]) + (float(out["水平距离（m）"]) + float(back["水平距离（m）"])) / float(r[5]) + descend / float(r[15])
            seconds = flight + float(r[10]) + len(ids) * float(r[11]) + float(r[12]) + len(ids) * float(r[13])
            assert mass <= float(r[3]) + 1e-8 and volume <= float(r[4]) + 1e-8
            assert actual <= (1 - reserve / 100) * float(r[8]) + 1e-8, (label, area, model, "SOC越限")
            for key, expected in (("箱数", len(ids)), ("质量_kg", mass), ("体积_m3", volume), ("能耗_kwh", actual), ("作业时间_s", seconds), ("返航SOC", 1 - actual / float(r[8]))):
                assert math.isclose(float(row[key]), expected, abs_tol=1e-8), (label, area, key, row[key], expected)
            assigned.extend(ids)
            totals[area] += 1
        assert Counter(assigned) == Counter(boxes.keys()), (label, "逐箱覆盖失败")
        assert len(totals) == 15
        return sum(float(row["能耗_kwh"]) for row in rows), sum(float(row["作业时间_s"]) for row in rows)

    for row in caps:
        area, model = row["服务区"], row["机型"]
        max_mass = float(drones[model][3])
        limit = 0.8 * float(drones[model][8])
        q = float(row["最大安全载荷_kg"])
        assert 0 <= q <= max_mass and energy(model, area, q) <= limit + 1e-8
        if row["状态"] == "energy_limited":
            assert abs(energy(model, area, q) - limit) <= 1e-7
        else:
            assert row["状态"] == "structural_payload_limit" and q == max_mass
    e20, t20 = check_rows(batches, 20.0, "基准")
    assert len(batches) == summary["总架次数"] == 18
    assert math.isclose(e20, summary["总能耗_kWh"], abs_tol=1e-8)
    assert math.isclose(t20, summary["累计作业时间_s"], abs_tol=1e-8)
    for reserve in (10.0, 20.0, 25.0, 30.0):
        rows = [row for row in scenarios if float(row["返航余量_百分比"]) == reserve]
        e, t = check_rows(rows, reserve, f"余量{reserve}")
        record = next(row for row in sensitivity if float(row["返航余量_百分比"]) == reserve)
        assert len(rows) == int(record["全任务架次数"])
        assert math.isclose(e, float(record["全任务能耗_kWh"]), abs_tol=1e-8)
        assert math.isclose(t, float(record["全任务累计时间_s"]), abs_tol=1e-8)
    result = {"status": "PASS", "boxes": len(boxes), "areas": 15, "caps": len(caps),
              "base_trips": len(batches), "scenarios": 4, "scenario_trips": len(scenarios)}
    (OUT / "独立审计报告.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result, ensure_ascii=True))


if __name__ == "__main__":
    main()
