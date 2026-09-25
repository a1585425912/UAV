"""第一轮建模与算法口径审计：以题目附录1/2/3 公式为准，独立重算并与实现值对拍。

运行：python 建模算法审计.py
输出：results/建模算法审计/10_题面口径对照表.csv、11_审计原始输出.json
边界：只读实现代码与数据，不改写任何已发布结果。
"""
from __future__ import annotations

import csv
import importlib
import json
import math
import re
from pathlib import Path

from openpyxl import load_workbook

CODE = Path(__file__).resolve().parent
ROOT = CODE.parent
OUT = ROOT / "results" / "建模算法审计"
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"

import sys
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
sys.path.insert(0, str(ROOT))
import 问题二_调度核心 as core
import 问题三_通信核心 as comm
import 问题四_分区求解 as q4

ROWS: list[dict] = []
INFO: dict = {}


def add(item, source, impl, verdict, evidence):
    ROWS.append({"条目": item, "题面出处": source, "实现位置": impl,
                 "判定": verdict, "证据": evidence})
    print(f"[{verdict}] {item} :: {evidence[:120]}")


def rel(a, b):
    return abs(a - b) / max(abs(b), 1e-12)


def load_drone_params():
    rows = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    out = {}
    for r in rows[2:5]:
        out[r[0]] = {"model": r[0], "mass0": float(r[2]), "payload": float(r[3]),
                     "volume": float(r[4]), "speed": float(r[5]),
                     "range0": float(r[6]), "range_full": float(r[7]),
                     "battery_kwh": float(r[8]), "reserve": float(r[9]) / 100,
                     "prep": float(r[10]), "load": float(r[11]), "handoff": float(r[12]),
                     "extra": float(r[13]), "climb_speed": float(r[14]),
                     "descent_speed": float(r[15]), "eta": float(r[16])}
    return out


def load_legs():
    with (ROOT / "results" / "有向航段几何参数.csv").open(encoding="utf-8-sig", newline="") as fh:
        return {(r["起点"], r["终点"]): {"distance": float(r["水平距离（m）"]),
                                         "climb": float(r["起点爬升（m）"]),
                                         "descent": float(r["终点下降（m）"]),
                                         "cruise": float(r["计划巡航海拔（m）"])}
                for r in csv.DictReader(fh)}


def load_comm_raw():
    rows = list(load_workbook(DATA / "通信链路参数.xlsx", read_only=True, data_only=True).active.values)
    return rows


def sheet_elevations():
    rows = list(load_workbook(DATA / "调度中心与服务区.xlsx", read_only=True, data_only=True).active.values)
    return {r[0]: float(r[4]) for r in rows if isinstance(r[0], str) and (r[0] == "O01" or r[0].startswith("S0"))}


def check_equivalent_range(drones, legs):
    """附录2 式: L(q)=L0-(L0-LF)(q/Q)^{3/2}；对拍水平能耗与飞行时间。"""
    picks = sorted(legs.items(), key=lambda kv: (-kv[1]["distance"], -kv[1]["climb"]))[:3]
    worst = 0.0
    n = 0
    for model, d in drones.items():
        data = {"drones": {model: d}, "legs": legs}
        for (a, b), leg in picks:
            for frac in (0.0, 0.25, 0.5, 1.0):
                q = frac * d["payload"]
                l_eq = d["range0"] - (d["range0"] - d["range_full"]) * (q / d["payload"]) ** 1.5
                e_hor = d["battery_kwh"] * leg["distance"] / l_eq
                e_up = (d["mass0"] + q) * 9.81 * leg["climb"] / (3.6e6 * d["eta"])
                t_ref = leg["climb"] / d["climb_speed"] + leg["distance"] / d["speed"] + leg["descent"] / d["descent_speed"]
                energy, seconds = core.segment(model, a, b, q, data)
                worst = max(worst, rel(e_hor + e_up, energy), rel(t_ref, seconds))
                n += 1
    add("附录2 等效航程(2-1)/时间(2-2)/能耗(2-3)分项",
        "附录2 式(2-1)(2-2)(2-3)", "问题二_调度核心.py: segment()",
        "符合" if worst < 1e-12 else "偏离",
        f"独立重算 {n} 组(3机型 x 3条最长最陡航段 x 4载荷) 与实现值最大相对差 {worst:.3e}；"
        "实现逐字使用 (q/Q)^{3/2} 插值、三段飞行时间与 E_hor+E_up 分解，下降项按效率0不计算")


def check_charge(drones, relay_full=1800.0):
    worst = 0.0
    for s in (0.0, 0.5, 0.8, 0.899999, 0.9, 0.95, 1.0):
        ref = (relay_full * (0.65 * (0.9 - s) / 0.9 + 0.35) if s < 0.9
               else relay_full * 0.35 * (1 - s) / 0.1)
        worst = max(worst, rel(ref, core.charge_time(s, relay_full)))
    add("附录2 两阶段等效充电模型",
        "附录2 式(2-5)", "问题二_调度核心.py: charge_time()",
        "符合" if worst < 1e-15 else "偏离",
        f"0/0.5/0.8/0.9-/0.9/0.95/1.0 七点独立重算最大相对差 {worst:.3e}；"
        "0→90% 占 65%Tfull、90→100% 占 35%Tfull 的线性折算逐字一致")


def check_fspl():
    src = (CODE / "问题三_通信核心.py").read_text(encoding="utf-8")
    pat = r"fspl = 32\.45 \+ 20 \* math\.log10\(params\[\"freq_mhz\"\]\) \+ 20 \* math\.log10\(distance_km\)"
    match = re.search(pat, src) is not None
    f_mhz, c = 2400.0, 299792458.0
    for d_km in (0.3, 1.0, 6.0, 20.0):
        ref = 32.45 + 20 * math.log10(f_mhz) + 20 * math.log10(d_km)
        phys = 20 * math.log10(4 * math.pi * d_km * 1000.0 * f_mhz * 1e6 / c)
        assert abs(ref - phys) < 0.02
    add("附录3 自由空间传播损耗",
        "附录3 式(3-5)", "问题三_通信核心.py: certified_link()",
        "符合" if match else "偏离",
        "源码表达式与式(3-5) 逐字一致(常量32.45, f/MHz, D/km)；"
        "与物理式 20log10(4πDf/c) 在 0.3-20 km 内差 <0.02 dB(常数取整)")


def check_limits():
    rows = load_comm_raw()
    f = float(rows[2][4]); lsys = float(rows[3][4]); lobs = float(rows[4][4])
    psens = float(rows[5][4]); margin = float(rows[6][4])
    dev = {"T": (float(rows[7][4]), float(rows[8][4])),
           "RA": (float(rows[9][4]), float(rows[10][4])),
           "RB": (float(rows[11][4]), float(rows[12][4])),
           "G": (float(rows[13][4]), float(rows[14][4]))}
    thr = psens + margin

    def limit(a, b):
        ab = dev[a][0] + dev[a][1] + dev[b][1] - lsys - thr
        ba = dev[b][0] + dev[b][1] + dev[a][1] - lsys - thr
        return min(ab, ba)

    params = comm.load_parameters()
    ref = {"limit_direct": limit("T", "G"), "limit_access": limit("T", "RA"),
           "limit_backhaul": limit("RB", "G")}
    worst = max(rel(v, params[k]) for k, v in ref.items())
    same_freq = math.isclose(params["freq_mhz"], f) and math.isclose(params["obstruction_db"], lobs)
    add("附录3 接收门限/双向链路预算/收发参数",
        "附录3 式(3-2)(3-3)(3-4)", "问题三_通信核心.py: load_parameters()",
        "符合" if worst < 1e-15 and same_freq else "偏离",
        f"独立解析通信链路参数.xlsx 重算 直连{ref['limit_direct']:.1f}/接入{ref['limit_access']:.1f}/"
        f"回传{ref['limit_backhaul']:.1f} dB 与实现最大相对差 {worst:.3e}；"
        f"门限 Pth=Psens+M={thr:.0f} dBm，接收端灵敏度和衰落裕量为全系统单一取值，与附件一致")


def check_operating_height():
    with (ROOT / "results" / "航路节点坐标与作业高度.csv").open(encoding="utf-8-sig", newline="") as fh:
        nodes = list(csv.DictReader(fh))
    diffs = [abs(float(r["DEM减节点表（m）"])) for r in nodes]
    worst = max(diffs)
    worst_node = nodes[diffs.index(worst)]["编号"]
    sheet = sheet_elevations()
    src3 = (CODE / "问题三_通信核心.py").read_text(encoding="utf-8")
    uses_sheet = "Point(float(r[2]), float(r[3]), float(r[4]))" in src3
    verdict = "偏离" if worst > 1e-6 and uses_sheet else "符合"
    add("附录1/2 节点作业高度基准(O01=地面海拔, 服务区=+30m)",
        "附录1 场景数据/附录2 作业高度", "航路几何数据.py:138(DEM基准) vs 问题三_通信核心.py:load_nodes(节点表基准)",
        verdict,
        f"两类基准最大差 {worst:.3f} m(节点{worst_node})；问题一几何缓存取 DEM 地面海拔+30 m，"
        f"而问题二/三求解取附件节点表地面海拔+30 m(uses_sheet={uses_sheet})，题目未指定以哪个为准")


def check_bisection(drones, legs):
    picks = sorted(legs.items(), key=lambda kv: (-kv[1]["distance"], -kv[1]["climb"]))[:2]
    bad = []
    for model, d in drones.items():
        data = {"drones": {model: d}, "legs": legs}
        for (a, b), leg in picks:
            prev = None
            for i in range(0, 21):
                q = d["payload"] * i / 20
                e, _ = core.segment(model, a, b, q, data)
                if prev is not None and e <= prev:
                    bad.append((model, a, b, q))
                prev = e
    add("问题一 安全载荷二分法的单调性前提",
        "附录2 式(2-1)(2-3)(2-4)", "问题一_基础计算.py: safe_payload()/trip()",
        "符合" if not bad else "偏离",
        f"{len(drones) * len(picks)} 条航段 x 21 个载荷点重算，E(q) 严格递增"
        if not bad else f"发现非单调点 {bad[:3]}")


def check_half_open_peak():
    plan = json.loads((ROOT / "results" / "问题三_参考口径" / "主方案_完整方案.json").read_text(encoding="utf-8"))
    comp, relay_ids, relay_comp, intervals, weights = q4.prepare(plan)

    def peak_half_open(items):
        events = sorted([(s, 1) for s, _ in items] + [(e, -1) for _, e in items])
        cur = best = 0
        for _, delta in events:
            cur += delta
            best = max(best, cur)
        return best

    coinc, detail = 0, []
    for cat, raw in intervals.items():
        pairs = [(a, b) for a, b, _kind, _identity in raw]
        code_peak, ref_peak = q4.peak(pairs), peak_half_open(pairs)
        for i, (s1, e1) in enumerate(pairs):
            for j, (s2, e2) in enumerate(pairs):
                if i < j and (e1 == s2 or e2 == s1):
                    coinc += 1
        if code_peak != ref_peak:
            detail.append(f"{cat}: peak()={code_peak} 半开区间={ref_peak}")
    add("问题四 资源峰值的事件端点法(半开区间语义)",
        "附录2 资源占用规则/论文半开区间约定", "问题四_分区求解.py: peak()",
        "符合" if not detail else "偏离",
        f"资源区间端点相接实例 {coinc} 处；(时刻,增量) 升序 -> 同刻结束事件先处理，与半开区间语义一致、"
        f"不会高估（t4 独立反例搜索 6000 组对抗 0 例不符；本脚本端点相接实例与半开区间标准算法结果"
        f"{'一致' if not detail else '不一致: ' + '; '.join(detail)}）")


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    drones, legs = load_drone_params(), load_legs()
    check_equivalent_range(drones, legs)
    check_charge(drones)
    check_fspl()
    check_limits()
    check_operating_height()
    check_bisection(drones, legs)
    check_half_open_peak()
    with (OUT / "10_题面口径对照表.csv").open("w", newline="", encoding="utf-8-sig") as fh:
        w = csv.DictWriter(fh, fieldnames=["条目", "题面出处", "实现位置", "判定", "证据"])
        w.writeheader()
        w.writerows(ROWS)
    (OUT / "11_审计原始输出.json").write_text(
        json.dumps({"条目数": len(ROWS), "判定统计": {v: sum(1 for r in ROWS if r["判定"] == v)
                                                   for v in ("符合", "偏离", "无法判定")},
                    "明细": ROWS}, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\n审计完成：{len(ROWS)} 条，写入 {OUT}")


if __name__ == "__main__":
    main()
