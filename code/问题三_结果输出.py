"""问题三中文 CSV/JSON 与官方工作簿 Q2/Q3 工作表。"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import load_workbook

from 问题二_调度核心 import load_data


ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "results" / "问题三_参考口径"
SOURCE_BOOK = ROOT / "results" / "问题二_参考口径" / "结果提交_主方案.xlsx"


def write_csv(path, fields, rows):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def save(result, stem="主方案"):
    OUT.mkdir(parents=True, exist_ok=True)
    data = load_data()
    trips = []
    for r in result["transport"]:
        trips.append({"架次编号": r["id"], "无人机编号": r["uav"], "机型编号": r["model"],
                      "电池编号": r["battery"], "开始时刻_s": r["start_s"],
                      "访问服务区顺序": "-".join(s["area"] for s in r["stops"]),
                      "返回O01时刻_s": r["return_s"], "架次能耗_kWh": r["energy_kwh"],
                      "返航SOC": r["soc"], "电池充满时刻_s": r["charge_end_s"],
                      "总质量_kg": r["mass_kg"], "总体积_m3": r["volume_m3"],
                      "货箱编号列表": "、".join(bid for stop in r["stops"] for bid in stop["ids"])})
    write_csv(OUT / f"{stem}_运输架次.csv", list(trips[0]), trips)
    deliveries = []
    for bid, when in sorted(result["deliveries"].items()):
        box = data["boxes"][bid]
        trip = next(r["id"] for r in result["transport"]
                    if any(bid in stop["ids"] for stop in r["stops"]))
        deliveries.append({"货箱编号": bid, "架次编号": trip, "服务区编号": box["area"],
                           "交付完成时刻_s": when, "期望时刻_s": box["due"],
                           "硬截止_s": box["hard"] if box["hard"] is not None else "",
                           "加权延误": box["weight"] * max(0, when - box["due"])})
    write_csv(OUT / f"{stem}_逐箱交付.csv", list(deliveries[0]), deliveries)
    relays = []
    for r in result["relays"]:
        relays.append({"中继架次编号": r["id"], "中继无人机编号": r["relay_id"],
                       "能源组件编号": r["energy_id"], "开始时刻_s": r["depart_s"],
                       "悬停经度": r["lon"], "悬停纬度": r["lat"], "悬停海拔_m": r["hover_alt_m"],
                       "建链完成时刻_s": r["link_ready_s"], "服务结束时刻_s": r["service_end_s"],
                       "返回O01时刻_s": r["return_s"], "架次能耗_kWh": r["energy_kwh"],
                       "返航SOC": r["soc"], "中继机可再用时刻_s": r["relay_free_s"],
                       "能源组件充满时刻_s": r["energy_free_s"]})
    write_csv(OUT / f"{stem}_中继架次.csv", list(relays[0]), relays)
    comm = [{"运输架次编号": r["trip"], "通信阶段": r["phase"],
             "开始时刻_s": r["start_s"], "结束时刻_s": r["end_s"],
             "保障方式": r["status"], "中继架次编号": r["relay_id"]} for r in result["communications"]]
    write_csv(OUT / f"{stem}_通信保障.csv", list(comm[0]), comm)
    full = {"metrics": result["metrics"], "transport": result["transport"],
            "relays": result["relays"], "deliveries": result["deliveries"],
            "communications": result["communications"]}
    (OUT / f"{stem}_完整方案.json").write_text(json.dumps(full, ensure_ascii=False, indent=2), encoding="utf-8")
    book = load_workbook(SOURCE_BOOK)
    def replace(sheet, rows, fields):
        ws = book[sheet]
        if ws.max_row > 1:
            ws.delete_rows(2, ws.max_row - 1)
        for row in rows:
            ws.append([row[field] for field in fields])
    replace("Q2_运输架次", trips, list(trips[0])[:8])
    replace("Q2_逐箱交付", deliveries, list(deliveries[0])[:4])
    replace("Q3_中继架次", relays, list(relays[0])[:11])
    replace("Q3_通信保障", comm, list(comm[0]))
    output = OUT / f"结果提交_{stem}.xlsx"
    book.save(output)
    print("已写入问题三", stem, result["metrics"], output)


if __name__ == "__main__":
    raise SystemExit("由 问题三_求解.py 调用")
