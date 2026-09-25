"""固定问题三全部架次与通信关系，用整数规划求 2/3 组资源配置。"""
from __future__ import annotations

import csv
import json
from collections import defaultdict
from itertools import combinations
from pathlib import Path

import numpy as np
from scipy.optimize import Bounds, LinearConstraint, milp
from scipy.sparse import coo_matrix
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
SOURCE = ROOT / "results" / "问题三_参考口径" / "主方案_完整方案.json"
OUT = ROOT / "results" / "问题四_参考口径"
CATEGORIES = ("U_A", "U_B", "U_C", "B_A", "B_B", "B_C", "R", "RB")
STOCK = {"U_A": 4, "U_B": 2, "U_C": 2, "B_A": 6, "B_B": 4, "B_C": 4, "R": 2, "RB": 6}


def components(plan):
    areas = sorted({stop["area"] for trip in plan["transport"] for stop in trip["stops"]})
    parent = {a: a for a in areas}
    def root(a):
        while parent[a] != a:
            a = parent[a]
        return a
    for trip in plan["transport"]:
        first = trip["stops"][0]["area"]
        for stop in trip["stops"][1:]:
            parent[root(stop["area"])] = root(first)
    groups = defaultdict(list)
    for a in areas:
        groups[root(a)].append(a)
    result = sorted((tuple(v) for v in groups.values()), key=lambda x: x[0])
    assert len(areas) == 15 and sum(map(len, result)) == 15
    return result


def prepare(plan):
    comp = components(plan)
    locate = {area: ci for ci, group in enumerate(comp) for area in group}
    trips = plan["transport"]
    relay_ids = sorted(r["id"] for r in plan["relays"])
    uses = defaultdict(set)
    for row in plan["communications"]:
        if row["relay_id"]:
            uses[row["trip"]].add(row["relay_id"])
    trip_comp = {trip["id"]: locate[trip["stops"][0]["area"]] for trip in trips}
    relay_comp = {rid: set() for rid in relay_ids}
    for trip_id, assigned in uses.items():
        for rid in assigned:
            relay_comp[rid].add(trip_comp[trip_id])
    intervals = {cat: [] for cat in CATEGORIES}
    weights = np.zeros(len(comp))
    for trip in trips:
        ci, model = trip_comp[trip["id"]], trip["model"]
        intervals[f"U_{model}"].append((trip["start_s"], trip["return_s"], "x", ci))
        intervals[f"B_{model}"].append((trip["start_s"], trip["charge_end_s"], "x", ci))
        weights[ci] += trip["return_s"] - trip["start_s"]
    for relay in plan["relays"]:
        rid = relay["id"]
        intervals["R"].append((relay["depart_s"], relay["relay_free_s"], "y", rid))
        intervals["RB"].append((relay["depart_s"], relay["energy_free_s"], "y", rid))
    return comp, relay_ids, relay_comp, intervals, weights


def peak(items):
    events = sorted([(start, 1) for start, end in items] + [(end, -1) for start, end in items],
                    key=lambda x: (x[0], x[1]))
    current = best = 0
    for _, delta in events:
        current += delta
        best = max(best, current)
    return best


def assess(assignment, prepared):
    comp, relay_ids, relay_comp, intervals, weights = prepared
    k = len(assignment)
    needs = []
    work = []
    for group in assignment:
        group_set = set(group)
        work.append(float(sum(weights[ci] for ci in group_set)))
        row = {}
        for category in CATEGORIES:
            taken = [(a, b) for a, b, kind, identity in intervals[category]
                     if (identity in group_set if kind == "x"
                         else bool(relay_comp[identity] & group_set))]
            row[category] = peak(taken)
        needs.append(row)
    total = {cat: sum(r[cat] for r in needs) for cat in CATEGORIES}
    gaps = {cat: max(0, total[cat] - STOCK[cat]) for cat in CATEGORIES}
    mean = sum(work) / k
    cv = (sum((v - mean) ** 2 for v in work) / k) ** .5 / mean
    return {"needs": needs, "total": total, "shortfall": gaps,
            "shortfall_total": sum(gaps.values()), "work_s": work, "cv": cv}


def solve(k, prepared, gap_allowance=0):
    comp, relay_ids, relay_comp, intervals, weights = prepared
    nc = len(comp)
    names = []
    def var(name):
        names.append(name)
        return len(names) - 1
    x = {(ci, group): var(("x", ci, group)) for ci in range(nc) for group in range(k)}
    y = {(rid, group): var(("y", rid, group)) for rid in relay_ids for group in range(k)}
    n = {(cat, group): var(("n", cat, group)) for cat in CATEGORIES for group in range(k)}
    delta = {cat: var(("delta", cat)) for cat in CATEGORIES}
    z = {group: var(("z", group)) for group in range(k)}
    count = len(names)
    integrality = np.zeros(count)
    lower = np.zeros(count)
    upper = np.full(count, np.inf)
    for index in list(x.values()) + list(y.values()):
        integrality[index] = 1
        upper[index] = 1
    for index in list(n.values()) + list(delta.values()):
        integrality[index] = 1
        upper[index] = 50
    rows, cols, vals, lows, highs = [], [], [], [], []
    def add(coef, lo=-np.inf, hi=np.inf):
        ri = len(lows)
        for col, value in coef.items():
            if value:
                rows.append(ri); cols.append(col); vals.append(value)
        lows.append(lo); highs.append(hi)
    for ci in range(nc):
        add({x[ci, group]: 1 for group in range(k)}, 1, 1)
    for group in range(k):
        add({x[ci, group]: 1 for ci in range(nc)}, 1)
    # 同一中继架次在任何有受保障运输的组中，都必须复制配置。
    for rid in relay_ids:
        for group in range(k):
            for ci in relay_comp[rid]:
                add({y[rid, group]: 1, x[ci, group]: -1}, 0)
            add({y[rid, group]: 1, **{x[ci, group]: -1 for ci in relay_comp[rid]}}, hi=0)
    for cat in CATEGORIES:
        active_times = sorted({start for start, _, _, _ in intervals[cat]})
        for group in range(k):
            for time in active_times:
                coef = {n[cat, group]: 1}
                for start, end, kind, identity in intervals[cat]:
                    if start <= time < end:
                        col = x[identity, group] if kind == "x" else y[identity, group]
                        coef[col] = coef.get(col, 0) - 1
                add(coef, 0)
    for cat in CATEGORIES:
        add({delta[cat]: 1, **{n[cat, group]: -1 for group in range(k)}}, -STOCK[cat])
    # W_k 只可能是不可分单元工作量的一个子集和；在全部可能 W_k 处加切线，
    # 整数可行点上 z_k=W_k^2 可精确实现，而无须枚举分区。
    hours = weights / 3600
    possible = {0.0}
    for value in hours:
        possible |= {old + float(value) for old in list(possible)}
    for group in range(k):
        for w0 in possible:
            coef = {z[group]: 1}
            for ci in range(nc):
                coef[x[ci, group]] = -2 * w0 * hours[ci]
            add(coef, -w0 * w0)
    matrix = coo_matrix((vals, (rows, cols)), shape=(len(lows), count)).tocsr()
    base_constraints = LinearConstraint(matrix, lows, highs)
    bounds = Bounds(lower, upper)
    def run(objective, extra=()):
        result = milp(objective, integrality=integrality, bounds=bounds,
                      constraints=[base_constraints, *extra],
                      options={"time_limit": 120, "mip_rel_gap": 1e-8})
        if result.x is None or result.status != 0:
            raise RuntimeError(f"K={k} 分区整数规划失败: {result.message}")
        return result
    def fixed_sum(indices, value):
        row = coo_matrix((np.ones(len(indices)), (np.zeros(len(indices)), indices)),
                         shape=(1, count)).tocsr()
        return LinearConstraint(row, value - 1e-6, value + 1e-6)
    deficit_cols = list(delta.values())
    resource_cols = list(n.values())
    c = np.zeros(count); c[deficit_cols] = 1
    first = run(c)
    minimum_gap = round(sum(first.x[col] for col in deficit_cols))
    gap_bound = fixed_sum(deficit_cols, minimum_gap) if gap_allowance == 0 else None
    if gap_allowance:
        row = coo_matrix((np.ones(len(deficit_cols)), (np.zeros(len(deficit_cols)), deficit_cols)),
                         shape=(1, count)).tocsr()
        gap_bound = LinearConstraint(row, 0, minimum_gap + gap_allowance + 1e-6)
    c = np.zeros(count); c[resource_cols] = 1
    second = run(c, [gap_bound])
    minimum_resources = round(sum(second.x[col] for col in resource_cols))
    c = np.zeros(count); c[list(z.values())] = 1
    additional = [gap_bound]
    if gap_allowance == 0:
        additional.append(fixed_sum(resource_cols, minimum_resources))
    final = run(c, additional)
    assignment = [tuple(ci for ci in range(nc) if final.x[x[ci, group]] > .5)
                  for group in range(k)]
    facts = assess(assignment, prepared)
    assert facts["shortfall_total"] <= minimum_gap + gap_allowance
    if gap_allowance == 0:
        assert facts["shortfall_total"] == minimum_gap
        assert sum(facts["total"].values()) == minimum_resources
    return assignment, facts, {"milp_status": final.status, "mip_gap": final.mip_gap,
                               "minimum_shortfall": minimum_gap,
                               "minimum_resources_given_shortfall": minimum_resources,
                               "binary_components": nc, "constraint_rows": len(lows)}


def save(k, name, assignment, facts, meta, prepared):
    comp = prepared[0]
    groups = [sorted(a for ci in indices for a in comp[ci]) for indices in assignment]
    assert sorted(a for group in groups for a in group) == sorted(a for c in comp for a in c)
    output = {"K": k, "方案": name, "任务组": groups, "不可分单元": [list(c) for c in comp],
              "资源需求": facts["needs"], "资源总需求": facts["total"], "库存": STOCK,
              "资源缺口": facts["shortfall"], "缺口总数": facts["shortfall_total"],
              "库存剩余": {cat: max(0, STOCK[cat] - facts["total"][cat]) for cat in STOCK},
              "各组运输作业量_s": facts["work_s"], "作业量CV": facts["cv"],
              "整数规划": meta}
    (OUT / f"K{k}_{name}.json").write_text(json.dumps(output, ensure_ascii=False, indent=2), encoding="utf-8")
    return output


def main():
    plan = json.loads(SOURCE.read_text(encoding="utf-8"))
    prepared = prepare(plan)
    OUT.mkdir(parents=True, exist_ok=True)
    answers = []
    for k in (2, 3):
        for name, allowance in (("缺口优先", 0), ("兼顾均衡", 4)):
            assignment, facts, meta = solve(k, prepared, allowance)
            answers.append(save(k, name, assignment, facts, meta, prepared))
            print(k, name, "缺口", facts["shortfall_total"], "CV", round(facts["cv"], 4),
                  "组", answers[-1]["任务组"])
    with (OUT / "问题四方案比较.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=["K", "方案", "缺口总数", "作业量CV", "资源总数", "任务组"])
        writer.writeheader()
        for row in answers:
            writer.writerow({key: value for key, value in {
                "K": row["K"], "方案": row["方案"], "缺口总数": row["缺口总数"],
                "作业量CV": row["作业量CV"], "资源总数": sum(row["资源总需求"].values()),
                "任务组": " / ".join("、".join(group) for group in row["任务组"])}.items()})
    primary = [a for a in answers if a["方案"] == "缺口优先"]
    detail = []
    for plan in primary:
        for group_index, (areas, need) in enumerate(zip(plan["任务组"], plan["资源需求"]), 1):
            detail.append({"K": plan["K"], "任务组编号": f"G{group_index}",
                           "服务区列表": "、".join(areas), **need,
                           "运输作业量_s": plan["各组运输作业量_s"][group_index - 1]})
    with (OUT / "问题四主方案逐组资源.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(detail[0]))
        writer.writeheader(); writer.writerows(detail)
    source_book = ROOT / "results" / "问题三_参考口径" / "结果提交_主方案.xlsx"
    book = load_workbook(source_book)
    ws = book["Q4_分区配置"]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    fields = ["K", "任务组编号", "服务区列表", *CATEGORIES]
    for row in detail:
        ws.append([row[field] for field in fields])
    book.save(OUT / "结果提交_问题三四主方案.xlsx")


if __name__ == "__main__":
    main()
