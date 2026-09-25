"""问题一直接指派 MILP 所用的输入、航段能耗和安全载荷计算。"""
from __future__ import annotations

import csv
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
G0 = 9.81
EPS = 1e-10


def read_csv(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def load_inputs():
    drone_rows = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    drones = {}
    for row in drone_rows[2:5]:
        assert row[0] in ("A", "B", "C")
        fields = ("mass0", "payload", "volume", "speed", "range0", "range_full", "energy",
                  "reserve", "prep", "load", "handoff", "extra", "climb_speed", "descent_speed", "eta")
        drones[row[0]] = dict(zip(fields, map(float, (*row[2:10], *row[10:17]))))
    assert set(drones) == {"A", "B", "C"}
    book = load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True)
    raw = list(book.worksheets[1].values)
    assert raw[0][:5] == ("货箱编号", "服务区编号", "物资类型", "单箱质量（kg）", "单箱体积（m³）")
    assert len(raw) == 81 and raw[-1][0] == "S015-FOD-01"
    boxes = defaultdict(list)
    for row in raw[1:]:
        assert row[0] and row[1] and float(row[3]) > 0 and float(row[4]) > 0
        boxes[row[1]].append({"id": row[0], "mass": float(row[3]),
                              "volume": float(row[4]), "type": row[2]})
    assert len({box["id"] for group in boxes.values() for box in group}) == 80
    assert len(boxes) == 15
    routes = {row["服务区"]: row for row in read_csv(ROOT / "results" / "单服务区往返几何参数.csv")}
    assert set(routes) == set(boxes)
    for area, route in routes.items():
        assert math.isclose(float(route["去程水平距离（m）"]), float(route["返程水平距离（m）"]), abs_tol=1e-7), area
        assert math.isclose(float(route["去程巡航海拔（m）"]), float(route["返程巡航海拔（m）"]), abs_tol=1e-7), area
        assert all(float(route[key]) >= 0 for key in
                   ("去程爬升（m）", "返程爬升（m）", "去程下降（m）", "返程下降（m）"))
    return drones, boxes, routes


def trip(drone, route, mass):
    payload = drone["payload"]
    assert 0 <= mass <= payload + EPS
    distance = float(route["去程水平距离（m）"])
    outward_climb = float(route["去程爬升（m）"])
    return_climb = float(route["返程爬升（m）"])
    loaded_range = drone["range0"] - (drone["range0"] - drone["range_full"]) * (mass / payload) ** 1.5
    parts = {
        "outbound_horizontal_kwh": drone["energy"] * distance / loaded_range,
        "outbound_climb_kwh": (drone["mass0"] + mass) * G0 * outward_climb / (3.6e6 * drone["eta"]),
        "return_horizontal_kwh": drone["energy"] * distance / drone["range0"],
        "return_climb_kwh": drone["mass0"] * G0 * return_climb / (3.6e6 * drone["eta"]),
    }
    parts["total_kwh"] = sum(parts.values())
    parts["return_soc"] = 1 - parts["total_kwh"] / drone["energy"]
    return parts


def operation_time(drone, route, count):
    distance = float(route["去程水平距离（m）"])
    climb = float(route["去程爬升（m）"]) + float(route["返程爬升（m）"])
    descend = float(route["去程下降（m）"]) + float(route["返程下降（m）"])
    flight = climb / drone["climb_speed"] + 2 * distance / drone["speed"] + descend / drone["descent_speed"]
    return flight + drone["prep"] + count * drone["load"] + drone["handoff"] + count * drone["extra"]


def safe_payload(drone, route):
    limit = (1 - drone["reserve"] / 100) * drone["energy"]
    empty = trip(drone, route, 0)["total_kwh"]
    full = trip(drone, route, drone["payload"])["total_kwh"]
    if empty > limit + EPS:
        return None, "infeasible_even_empty", empty, full
    if full <= limit:
        return drone["payload"], "structural_payload_limit", empty, full
    lower, upper = 0.0, drone["payload"]
    for _ in range(55):
        mid = (lower + upper) / 2
        if trip(drone, route, mid)["total_kwh"] <= limit:
            lower = mid
        else:
            upper = mid
    return lower, "energy_limited", empty, full
