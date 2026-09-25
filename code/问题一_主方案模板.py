"""把当前非枚举主方案写入官方结果模板的 Q1 工作表。"""
from __future__ import annotations

import csv
from copy import copy
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parent.parent
RESULT = ROOT / "results" / "问题一_非枚举整数规划"
DEST = RESULT / "结果提交_问题一主方案.xlsx"
OFFICIAL = ROOT / "结果提交模板.xlsx"
SOURCE = RESULT / "最优组批_逐架次.csv"


def main():
    with SOURCE.open(encoding="utf-8-sig", newline="") as stream:
        trips = list(csv.DictReader(stream))
    assert len(trips) == 18
    template = OFFICIAL if OFFICIAL.exists() else DEST
    if not template.exists():
        raise FileNotFoundError("需要原题结果提交模板，或仓库保存的 Q1 主方案工作簿")
    book = load_workbook(template)
    sheets = book.sheetnames[:]
    ws = book["Q1_单点组批"]
    columns = ("架次编号", "服务区编号", "机型编号", "货箱编号列表", "总质量（kg）",
               "总体积（m³）", "往返时间（s）", "架次能耗（kWh）", "返航SOC（%）")
    assert tuple(ws.cell(1, column).value for column in range(1, 10)) == columns
    model_style = [(copy(ws.cell(2, column)._style), copy(ws.cell(2, column).alignment))
                   for column in range(1, 10)]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    for number, trip in enumerate(trips, 1):
        values = (f"Q1-{number:02d}", trip["服务区"], trip["机型"], trip["货箱编号列表"],
                  float(trip["质量_kg"]), float(trip["体积_m3"]), float(trip["作业时间_s"]),
                  float(trip["能耗_kwh"]), 100 * float(trip["返航SOC"]))
        for column, (value, (cell_style, alignment)) in enumerate(zip(values, model_style), 1):
            cell = ws.cell(number + 1, column, value)
            cell._style = copy(cell_style)
            cell.alignment = copy(alignment)
    assert book.sheetnames == sheets
    book.save(DEST)
    verify = load_workbook(DEST, read_only=True, data_only=True)
    rows = list(verify["Q1_单点组批"].values)
    assert len(rows) == 19 and rows[1][0] == "Q1-01" and rows[-1][0] == "Q1-18"
    verify.close()
    print(f"已将非枚举主方案 18 架次写入：{DEST}")


if __name__ == "__main__":
    main()
