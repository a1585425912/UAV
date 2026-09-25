"""问题二多点运输物理模型、逐箱交付时刻与实体资源解码器。"""
from __future__ import annotations

import argparse
import csv
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
GEOMETRY = ROOT / "results" / "问题二_参考口径" / "有向航段几何参数.csv"
GRAVITY = 9.81
EPS = 1e-8


def read_geometry(path=GEOMETRY):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 240
    return {(r["起点"], r["终点"]): {
        "distance": float(r["水平距离（m）"]),
        "climb": float(r["起点爬升（m）"]),
        "descent": float(r["终点下降（m）"]),
        "cruise": float(r["计划巡航海拔（m）"]),
    } for r in rows}


def load_data():
    raw = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    assert raw[1][0] == "机型编号" and raw[2][0] == "A" and raw[4][0] == "C"
    drones = {}
    for r in raw[2:5]:
        drones[r[0]] = {
            "mass0": float(r[2]), "payload": float(r[3]), "volume": float(r[4]),
            "speed": float(r[5]), "range0": float(r[6]), "range_full": float(r[7]),
            "battery_kwh": float(r[8]), "reserve": float(r[9]) / 100,
            "prep": float(r[10]), "load": float(r[11]),
            "handoff": float(r[12]), "extra": float(r[13]),
            "climb_speed": float(r[14]), "descent_speed": float(r[15]),
            "eta": float(r[16]),
        }
    assert set(drones) == {"A", "B", "C"}
    uavs = {g: [] for g in drones}
    for r in raw[8:16]:
        assert r[0] and r[1] in drones and r[2] == "O01"
        uavs[r[1]].append(r[0])
    assert {g: len(v) for g, v in uavs.items()} == {"A": 4, "B": 2, "C": 2}
    batteries = {}
    for r in raw[19:22]:
        batteries[r[0]] = {"ids": [f"{r[0]}-B{i}" for i in range(1, int(r[1]) + 1)],
                            "full_charge_s": float(r[2])}
    assert {g: len(v["ids"]) for g, v in batteries.items()} == {"A": 6, "B": 4, "C": 4}
    workbook = load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True)
    raw_boxes = list(workbook.worksheets[1].values)
    assert raw_boxes[0][0] == "货箱编号" and len(raw_boxes) == 81
    boxes = {}
    for r in raw_boxes[1:]:
        is_first = r[5] == "是"
        limits = ([float(r[6])] if is_first and r[6] is not None else [])
        if r[2] == "医疗物资":
            limits.append(float(r[7]))
        boxes[r[0]] = {"area": r[1], "type": r[2], "mass": float(r[3]),
                       "volume": float(r[4]), "first": is_first,
                       "hard": min(limits) if limits else None,
                       "due": float(r[7]), "weight": float(r[8])}
    assert len(boxes) == 80 and len({v["area"] for v in boxes.values()}) == 15
    assert sum(b["first"] for b in boxes.values()) == 30
    assert Counter(b["hard"] for b in boxes.values()) == Counter({None: 49, 3600.0: 17, 7200.0: 7, 10800.0: 7})
    return {"drones": drones, "uavs": uavs, "batteries": batteries,
            "boxes": boxes, "legs": read_geometry()}


def charge_time(soc, full_charge_s):
    assert -EPS <= soc <= 1 + EPS
    soc = min(1.0, max(0.0, soc))
    if soc < .9:
        return full_charge_s * (.65 * (.9 - soc) / .9 + .35)
    return full_charge_s * .35 * (1 - soc) / .1


def segment(model, start, end, payload, data):
    drone, leg = data["drones"][model], data["legs"][start, end]
    if not (0 <= payload <= drone["payload"] + EPS):
        raise ValueError("航段载荷越界")
    q = min(payload, drone["payload"])
    equivalent_range = drone["range0"] - (drone["range0"] - drone["range_full"]) * (q / drone["payload"]) ** 1.5
    energy = (drone["battery_kwh"] * leg["distance"] / equivalent_range
              + (drone["mass0"] + q) * GRAVITY * leg["climb"] / (3.6e6 * drone["eta"]))
    seconds = (leg["climb"] / drone["climb_speed"]
               + leg["distance"] / drone["speed"]
               + leg["descent"] / drone["descent_speed"])
    return energy, seconds


def evaluate_sortie(sortie, data):
    model, stops = sortie["model"], sortie["stops"]
    if model not in data["drones"] or not stops:
        return None
    drone, boxes = data["drones"][model], data["boxes"]
    ids = [bid for stop in stops for bid in stop["ids"]]
    if not ids or len(ids) != len(set(ids)) or any(bid not in boxes for bid in ids):
        return None
    if any(not stop["ids"] or len({boxes[bid]["area"] for bid in stop["ids"]}) != 1
           or boxes[stop["ids"][0]]["area"] != stop["area"] for stop in stops):
        return None
    if len({stop["area"] for stop in stops}) != len(stops):
        return None
    remaining_mass = sum(boxes[bid]["mass"] for bid in ids)
    volume = sum(boxes[bid]["volume"] for bid in ids)
    if remaining_mass > drone["payload"] + EPS or volume > drone["volume"] + EPS:
        return None
    elapsed = drone["prep"] + len(ids) * drone["load"]
    energy, location, deliveries, legs = 0.0, "O01", {}, []
    for stop in stops:
        area = stop["area"]
        e, seconds = segment(model, location, area, remaining_mass, data)
        energy += e
        elapsed += seconds
        legs.append({"from": location, "to": area, "payload_kg": remaining_mass,
                     "energy_kwh": e, "flight_s": seconds})
        elapsed += drone["handoff"] + len(stop["ids"]) * drone["extra"]
        deliveries.update({bid: elapsed for bid in stop["ids"]})
        remaining_mass -= sum(boxes[bid]["mass"] for bid in stop["ids"])
        location = area
    if abs(remaining_mass) > 1e-7:
        raise AssertionError("末站后载荷未清零")
    e, seconds = segment(model, location, "O01", 0.0, data)
    energy += e
    elapsed += seconds
    legs.append({"from": location, "to": "O01", "payload_kg": 0.0,
                 "energy_kwh": e, "flight_s": seconds})
    soc = 1 - energy / drone["battery_kwh"]
    if soc < drone["reserve"] - EPS:
        return None
    return {"model": model, "stops": stops, "box_ids": ids, "mass_kg": sum(boxes[bid]["mass"] for bid in ids),
            "volume_m3": volume, "duration_s": elapsed, "energy_kwh": energy,
            "return_soc": soc, "delivery_offsets_s": deliveries, "legs": legs}


def decode(sorties, data, require_all=False):
    box_ids = [bid for s in sorties for stop in s["stops"] for bid in stop["ids"]]
    if len(box_ids) != len(set(box_ids)) or (require_all and Counter(box_ids) != Counter(data["boxes"].keys())):
        return None
    uav_free = {uid: 0.0 for group in data["uavs"].values() for uid in group}
    battery_free = {bid: 0.0 for group in data["batteries"].values() for bid in group["ids"]}
    result, delivery = [], {}
    for index, spec in enumerate(sorties, 1):
        ev = evaluate_sortie(spec, data)
        if ev is None:
            return None
        model = spec["model"]
        start, uid, bid = min((max(uav_free[u], battery_free[b]), u, b)
                              for u in data["uavs"][model]
                              for b in data["batteries"][model]["ids"])
        ret = start + ev["duration_s"]
        charged = ret + charge_time(ev["return_soc"], data["batteries"][model]["full_charge_s"])
        uav_free[uid], battery_free[bid] = ret, charged
        delivery.update({box: start + offset for box, offset in ev["delivery_offsets_s"].items()})
        result.append({"id": f"T{index:02d}", "model": model, "uav": uid, "battery": bid,
                       "start_s": start, "return_s": ret, "charge_end_s": charged,
                       "energy_kwh": ev["energy_kwh"], "return_soc": ev["return_soc"],
                       "stops": spec["stops"], "box_ids": ev["box_ids"], "mass_kg": ev["mass_kg"],
                       "volume_m3": ev["volume_m3"]})
    boxes = data["boxes"]
    hard_over = sum(max(0.0, delivery[bid] - boxes[bid]["hard"])
                    for bid in delivery if boxes[bid]["hard"] is not None)
    weighted_tardiness = sum(boxes[bid]["weight"] * max(0.0, t - boxes[bid]["due"])
                              for bid, t in delivery.items())
    metrics = {"hard_excess_s": hard_over, "weighted_tardiness": weighted_tardiness,
               "makespan_s": max((s["return_s"] for s in result), default=0.0),
               "energy_kwh": sum(s["energy_kwh"] for s in result), "sorties": len(result),
               "boxes": len(delivery)}
    return {"sorties": result, "deliveries": delivery, "metrics": metrics}


def smoke():
    data = load_data()
    boxes = data["boxes"]
    def med_water(area):
        return [next(bid for bid, b in boxes.items() if b["area"] == area and b["type"] == kind)
                for kind in ("医疗物资", "饮用水")]
    triple = {"model": "C", "stops": [{"area": area, "ids": med_water(area)}
                                      for area in ("S013", "S010", "S014")]}
    first = evaluate_sortie(triple, data)
    assert first is not None and len(first["box_ids"]) == 6
    plan = decode([triple, {"model": "A", "stops": [{"area": "S001", "ids": ["S001-MED-01"]}]},
                   {"model": "C", "stops": [{"area": "S006", "ids": med_water("S006")}]},
                   triple], data)
    assert plan is None  # 重复箱必须被拒绝
    plan = decode([triple, {"model": "A", "stops": [{"area": "S001", "ids": ["S001-MED-01"]}]},
                   {"model": "C", "stops": [{"area": "S006", "ids": med_water("S006")}]},
                   {"model": "C", "stops": [{"area": "S007", "ids": med_water("S007")}]},
                   ], data)
    assert plan is not None and len(plan["sorties"]) == 4
    for row in plan["sorties"]:
        assert row["return_soc"] >= data["drones"][row["model"]]["reserve"] - EPS
        assert row["charge_end_s"] > row["return_s"]
    print(json.dumps({"代表性三站路线能耗_kWh": first["energy_kwh"],
                      "参考值_kWh": 5.100, "充电示例_A_71.5%_s": charge_time(.715, 1800),
                      "四架次纵向切片": plan["metrics"],
                      "首架次": plan["sorties"][0]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--smoke", action="store_true")
    args = parser.parse_args()
    if args.smoke:
        smoke()
    else:
        parser.error("当前仅提供 --smoke；全量搜索由问题二求解脚本调用")
