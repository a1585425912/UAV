"""从原始货箱表、机型表、统一有向航段和问题一输出独立复算约束。"""
from __future__ import annotations

import csv
import json
import math
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
OUT = ROOT / "results" / "问题一_统一航段"


def rows(path):
    with path.open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def main():
    drones_raw = list(load_workbook(DATA / "运输无人机数据.xlsx", read_only=True, data_only=True).active.values)
    drones = {r[0]: r for r in drones_raw[2:5]}
    boxes_raw = list(load_workbook(DATA / "物资需求与配送时限.xlsx", read_only=True, data_only=True).worksheets[1].values)
    boxes = {r[0]: r for r in boxes_raw[1:]}
    legs = {(r["起点"], r["终点"]): r for r in rows(ROOT / "results" / "有向航段几何参数.csv")}
    caps = rows(OUT / "最大安全载荷_45组.csv")
    details = rows(OUT / "安全载荷能耗明细.csv")
    candidates = rows(OUT / "全部可行组批.csv")
    chosen = rows(OUT / "最优组批_逐架次.csv")
    summary = json.loads((OUT / "汇总.json").read_text(encoding="utf-8"))
    assert len(drones) == 3 and len(boxes) == 80 and len(legs) == 240
    assert len(caps) == len(details) == 45
    assert len({(r["服务区"], r["机型"]) for r in caps}) == 45
    candidate_keys = {(r["服务区"], r["子集掩码"], r["机型"]) for r in candidates}
    assert len(candidate_keys) == len(candidates)

    def energy(model, area, mass):
        r = drones[model]
        out = legs[("O01", area)]
        back = legs[(area, "O01")]
        Q, m0, L0, LF, E, eta = map(float, (r[3], r[2], r[6], r[7], r[8], r[16]))
        d0, d1 = float(out["水平距离（m）"]), float(back["水平距离（m）"])
        L = L0 - (L0 - LF) * (mass / Q) ** 1.5
        return (E * d0 / L + E * d1 / L0
                + 9.81 * ((m0 + mass) * float(out["起点爬升（m）"])
                          + m0 * float(back["起点爬升（m）"])) / (3.6e6 * eta))

    errors = []
    for row in details:
        area, model = row["服务区"], row["机型"]
        r = drones[model]
        Q, E = float(r[3]), float(r[8])
        limit = 0.8 * E
        if not math.isclose(float(row["空载能耗_kwh"]), energy(model, area, 0), abs_tol=1e-8):
            errors.append(f"{area}/{model}: 空载能耗差异")
        if not math.isclose(float(row["满载能耗_kwh"]), energy(model, area, Q), abs_tol=1e-8):
            errors.append(f"{area}/{model}: 满载能耗差异")
        if row["状态"] == "infeasible_even_empty":
            if energy(model, area, 0) <= limit + 1e-9:
                errors.append(f"{area}/{model}: 空载不可行标记错误")
        else:
            q = float(row["最大安全载荷_kg"])
            if not (0 <= q <= Q and energy(model, area, q) <= limit + 1e-8):
                errors.append(f"{area}/{model}: 安全载荷不可行")
            if row["状态"] == "structural_payload_limit" and q != Q:
                errors.append(f"{area}/{model}: 额定上限不等于Q")
            if row["状态"] == "energy_limited" and (Q - q < 1e-5 or abs(energy(model, area, q) - limit) > 1e-7):
                errors.append(f"{area}/{model}: 能量边界偏差")

    def check_batch(row, label):
        area, model = row["服务区"], row["机型"]
        r = drones[model]
        ids = row["货箱编号列表"].split(",")
        mass = sum(float(boxes[bid][3]) for bid in ids)
        volume = sum(float(boxes[bid][4]) for bid in ids)
        e = energy(model, area, mass)
        soc = 1 - e / float(r[8])
        if any(boxes[bid][1] != area for bid in ids):
            errors.append(f"{label}: 跨区货箱")
        if len(ids) != len(set(ids)):
            errors.append(f"{label}: 同架重复货箱")
        if not (mass <= float(r[3]) + 1e-9 and volume <= float(r[4]) + 1e-9 and e <= 0.8 * float(r[8]) + 1e-9):
            errors.append(f"{label}: 质量/体积/能耗越限")
        for k, actual in (("质量_kg", mass), ("体积_m3", volume), ("能耗_kwh", e), ("返航SOC", soc)):
            if not math.isclose(float(row[k]), actual, abs_tol=1e-8):
                errors.append(f"{label}: {k}不一致")
        outbound = legs[("O01", area)]
        backward = legs[(area, "O01")]
        fly = ((float(outbound["起点爬升（m）"]) + float(backward["起点爬升（m）"])) / float(r[14])
               + (float(outbound["水平距离（m）"]) + float(backward["水平距离（m）"])) / float(r[5])
               + (float(outbound["终点下降（m）"]) + float(backward["终点下降（m）"])) / float(r[15]))
        t = fly + float(r[10]) + len(ids) * float(r[11]) + float(r[12]) + len(ids) * float(r[13])
        if not math.isclose(float(row["作业时间_s"]), t, abs_tol=1e-7):
            errors.append(f"{label}: 作业时间不一致")
        return ids

    for j, row in enumerate(candidates):
        check_batch(row, f"候选{j+1}")
    allocated = []
    for j, row in enumerate(chosen):
        allocated.extend(check_batch(row, f"最优{j+1}"))
        if (row["服务区"], row["子集掩码"], row["机型"]) not in candidate_keys:
            errors.append(f"最优{j+1}: 不在候选库")
    if Counter(allocated) != Counter(boxes.keys()):
        errors.append("80箱覆盖或唯一性错误")
    if len(chosen) != summary["trips"] or len(candidates) != summary["candidate_rows"]:
        errors.append("汇总数量不一致")
    if not math.isclose(sum(float(r["能耗_kwh"]) for r in chosen), summary["energy_kwh"], abs_tol=1e-8):
        errors.append("汇总能耗不一致")
    if not math.isclose(sum(float(r["作业时间_s"]) for r in chosen), summary["operation_s"], abs_tol=1e-6):
        errors.append("汇总时间不一致")
    report = [f"状态: {'PASS' if not errors else 'FAIL'}", "输入: 原始无人机/逐箱xlsx、240条有向航段、问题一输出", f"载荷组合: {len(caps)}", f"候选组批: {len(candidates)}", f"最优架次: {len(chosen)}", f"货箱覆盖: {len(allocated)}/80", f"发现: {len(errors)}"] + errors[:100]
    (OUT / "独立审计报告.txt").write_text("\n".join(report) + "\n", encoding="utf-8")
    print("\n".join(report[:7]))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
