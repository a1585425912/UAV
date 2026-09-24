"""从原始工作簿和交付表独立复算问题三运输与问题四分区资源。"""
from __future__ import annotations

import csv
import json
import math
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image

from 问题二_独立审计 import close, inputs, leg_value, recharge
from 地理计算 import geodesic_distance


ROOT = Path(__file__).resolve().parent
Q3 = ROOT / "results" / "问题三_参考口径"
Q4 = ROOT / "results" / "问题四_参考口径"
STOCK = {"U_A": 4, "U_B": 2, "U_C": 2, "B_A": 6, "B_B": 4, "B_C": 4, "R": 2, "RB": 6}


def read(path):
    with path.open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def peak(intervals):
    events = sorted([(a, 1) for a, b in intervals] + [(b, -1) for a, b in intervals],
                    key=lambda e: (e[0], e[1]))
    active = best = 0
    for _, delta in events:
        active += delta
        best = max(best, active)
    return best


def exact_dem_max(image, lon0, lat0, lon1, lat1):
    """独立栅格边界穿越：列出直线进入的每个 DEM 像元。"""
    tie, size = image.tag_v2[33922], image.tag_v2[33550]
    c0, c1 = (lon0 - tie[3]) / size[0], (lon1 - tie[3]) / size[0]
    r0, r1 = (tie[4] - lat0) / size[1], (tie[4] - lat1) / size[1]
    ticks = [0.0, 1.0]
    for a, b in ((c0, c1), (r0, r1)):
        if abs(b - a) < 1e-12:
            continue
        for boundary in range(math.floor(min(a, b)) + 1, math.ceil(max(a, b))):
            t = (boundary - a) / (b - a)
            if 0 < t < 1:
                ticks.append(t)
    heights = []
    ticks = sorted(set(ticks))
    for left, right in zip(ticks, ticks[1:]):
        mid = (left + right) / 2
        col = math.floor(c0 + mid * (c1 - c0))
        row = math.floor(r0 + mid * (r1 - r0))
        heights.append(float(image.getpixel((col, row))))
    return max(heights)


def verify_q3(stem, source):
    models, uavs, batteries, boxes, legs = source
    trips = read(Q3 / f"{stem}_运输架次.csv")
    deliveries = read(Q3 / f"{stem}_逐箱交付.csv")
    relays = read(Q3 / f"{stem}_中继架次.csv")
    comm = read(Q3 / f"{stem}_通信保障.csv")
    assert len(trips) == 22 and len(deliveries) == 80 and len(relays) == 4
    assert {r["货箱编号"] for r in deliveries} == set(boxes)
    by_trip = defaultdict(list)
    for d in deliveries:
        by_trip[d["架次编号"]].append(d)
        assert boxes[d["货箱编号"]][1] == d["服务区编号"]
    assert set(by_trip) == {t["架次编号"] for t in trips}
    uav_iv, battery_iv, relay_iv, energy_iv = defaultdict(list), defaultdict(list), defaultdict(list), defaultdict(list)
    hard, weighted, transport_energy = 0.0, 0.0, 0.0
    relay_by_id = {r["中继架次编号"]: r for r in relays}
    for trip in trips:
        sid, uid, model_name, bid = (trip[k] for k in ("架次编号", "无人机编号", "机型编号", "电池编号"))
        model = models[model_name]
        assert uavs[uid] == model_name
        assert bid.startswith(model_name + "-B")
        assigned = by_trip[sid]
        ids = {d["货箱编号"] for d in assigned}
        assert set(trip["货箱编号列表"].split("、")) == ids
        mass = sum(float(boxes[i][3]) for i in ids)
        volume = sum(float(boxes[i][4]) for i in ids)
        close(mass, trip["总质量_kg"], (sid, "质量"))
        close(volume, trip["总体积_m3"], (sid, "体积"))
        assert mass <= float(model[3]) + 1e-8 and volume <= float(model[4]) + 1e-8
        areas = trip["访问服务区顺序"].split("-")
        assert len(areas) == len(set(areas)) and set(areas) == {d["服务区编号"] for d in assigned}
        start = float(trip["开始时刻_s"])
        elapsed = float(model[10]) + len(ids) * float(model[11])
        remaining, origin, used = mass, "O01", 0.0
        for area in areas:
            energy, seconds = leg_value(model, legs[origin, area], remaining)
            elapsed += seconds
            used += energy
            here = [d for d in assigned if d["服务区编号"] == area]
            elapsed += float(model[12]) + len(here) * float(model[13])
            for d in here:
                box = boxes[d["货箱编号"]]
                actual = start + elapsed
                close(actual, d["交付完成时刻_s"], (sid, d["货箱编号"], "送达"))
                deadline = []
                if box[5] == "是" and box[6] is not None:
                    deadline.append(float(box[6]))
                if box[2] == "医疗物资":
                    deadline.append(float(box[7]))
                if deadline:
                    hard += max(0, actual - min(deadline))
                weighted += float(box[8]) * max(0, actual - float(box[7]))
            remaining -= sum(float(boxes[d["货箱编号"]][3]) for d in here)
            origin = area
        assert abs(remaining) < 1e-7
        energy, seconds = leg_value(model, legs[origin, "O01"], 0)
        elapsed += seconds
        used += energy
        ret = start + elapsed
        soc = 1 - used / float(model[8])
        assert soc >= float(model[9]) / 100 - 1e-8
        charge_end = ret + recharge(soc, batteries[model_name][1])
        for key, expected in (("返回O01时刻_s", ret), ("架次能耗_kWh", used),
                              ("返航SOC", soc), ("电池充满时刻_s", charge_end)):
            close(trip[key], expected, (sid, key))
        uav_iv[uid].append((start, ret))
        battery_iv[bid].append((start, charge_end))
        transport_energy += used
    assert hard <= 1e-8 and weighted <= 1e-8
    for name, intervals in list(uav_iv.items()) + list(battery_iv.items()):
        ordered = sorted(intervals)
        assert all(a[1] <= b[0] + 1e-8 for a, b in zip(ordered, ordered[1:])), name
    # 逐段通讯表必须覆盖全部飞行与投送时间，且中继声明落在实际服务窗口内。
    comm_by_trip = defaultdict(list)
    for row in comm:
        assert row["保障方式"] in {"直连", "中继"}
        a, b = float(row["开始时刻_s"]), float(row["结束时刻_s"])
        assert b > a
        if row["保障方式"] == "中继":
            relay = relay_by_id[row["中继架次编号"]]
            assert float(relay["建链完成时刻_s"]) <= a + 1e-8
            assert b <= float(relay["服务结束时刻_s"]) + 1e-8
        else:
            assert not row["中继架次编号"]
        comm_by_trip[row["运输架次编号"]].append((a, b))
    assert set(comm_by_trip) == {t["架次编号"] for t in trips}
    for trip in trips:
        sid, model = trip["架次编号"], models[trip["机型编号"]]
        covered = sorted(comm_by_trip[sid])
        start = float(trip["开始时刻_s"])
        prep_end = start + float(model[10]) + len(by_trip[sid]) * float(model[11])
        close(covered[0][0], prep_end, (sid, "通信起点"))
        close(covered[-1][1], trip["返回O01时刻_s"], (sid, "通信终点"))
        for left, right in zip(covered, covered[1:]):
            close(left[1], right[0], (sid, "通信段连接"))
    # 资源身份、机身周转与能源组件充电另外从原始中继表检验。
    raw = list(load_workbook(ROOT / "数据" / "无人机应急物资运输基础数据" / "中继无人机数据.xlsx",
                             read_only=True, data_only=True).active.values)
    rmodel = raw[2]
    node_raw = list(load_workbook(ROOT / "数据" / "无人机应急物资运输基础数据" / "调度中心与服务区.xlsx",
                                  read_only=True, data_only=True).active.values)
    o_lon, o_lat, o_alt = map(float, node_raw[2][2:5])
    dem = Image.open(ROOT / "数据" / "镇龙乡及周边30米DEM.tif")
    full = json.loads((Q3 / f"{stem}_完整方案.json").read_text(encoding="utf-8"))
    raw_relay = {r["id"]: r for r in full["relays"]}
    used_relays, used_energies = set(), set()
    relay_energy = 0.0
    for r in relays:
        uid, eid = r["中继无人机编号"], r["能源组件编号"]
        assert uid in {"R01", "R02"} and eid in {f"R-B{i}" for i in range(1, 7)}
        used_relays.add(uid); used_energies.add(eid)
        start, ret = float(r["开始时刻_s"]), float(r["返回O01时刻_s"])
        relay = raw_relay[r["中继架次编号"]]
        lon, lat, alt = (float(r[key]) for key in ("悬停经度", "悬停纬度", "悬停海拔_m"))
        ground = dem.getpixel((math.floor((lon - dem.tag_v2[33922][3]) / dem.tag_v2[33550][0]),
                               math.floor((dem.tag_v2[33922][4] - lat) / dem.tag_v2[33550][1])))
        assert ground <= alt <= ground + float(rmodel[18]) + 1e-8
        cruise = float(relay["cruise_alt_m"])
        assert cruise >= exact_dem_max(dem, o_lon, o_lat, lon, lat) + 50 - 1e-8
        assert cruise >= alt - 1e-8
        distance = geodesic_distance(o_lon, o_lat, lon, lat)
        close(distance, relay["distance_m"], (r["中继架次编号"], "水平距离"))
        out_time = ((cruise - o_alt) / float(rmodel[12]) + distance / float(rmodel[5])
                    + (cruise - alt) / float(rmodel[13]))
        back_time = ((cruise - alt) / float(rmodel[12]) + distance / float(rmodel[5])
                     + (cruise - o_alt) / float(rmodel[13]))
        ready = start + float(rmodel[9]) + out_time + float(rmodel[10])
        service_end = float(r["服务结束时刻_s"])
        expected_ret = service_end + back_time
        transit = (2 * float(rmodel[6]) * distance / float(rmodel[5]) / 3600
                   + float(rmodel[4]) * 9.81 * ((cruise - o_alt) + (cruise - alt))
                   / (3.6e6 * float(rmodel[14])))
        service = (float(rmodel[16]) + float(rmodel[17])) * (
            service_end - (ready - float(rmodel[10]))) / 3600
        close(ready, r["建链完成时刻_s"], (r["中继架次编号"], "建链"))
        close(expected_ret, ret, (r["中继架次编号"], "返航"))
        close(transit + service, r["架次能耗_kWh"], (r["中继架次编号"], "能耗"))
        end_machine = ret + float(rmodel[11])
        soc = 1 - float(r["架次能耗_kWh"]) / float(rmodel[7])
        charge_end = ret + recharge(soc, float(raw[11][2]))
        close(r["中继机可再用时刻_s"], end_machine, (r["中继架次编号"], "周转"))
        close(r["能源组件充满时刻_s"], charge_end, (r["中继架次编号"], "充电"))
        close(r["返航SOC"], soc, (r["中继架次编号"], "SOC"))
        assert soc >= float(rmodel[8]) / 100 - 1e-8
        relay_iv[uid].append((start, end_machine))
        energy_iv[eid].append((start, charge_end))
        relay_energy += float(r["架次能耗_kWh"])
    for name, intervals in list(relay_iv.items()) + list(energy_iv.items()):
        ordered = sorted(intervals)
        assert all(a[1] <= b[0] + 1e-8 for a, b in zip(ordered, ordered[1:])), name
    metrics = full["metrics"]
    close(metrics["energy_kwh"], relay_energy + transport_energy, "总能耗")
    close(metrics["makespan_s"], max([float(t["返回O01时刻_s"]) for t in trips]
                                    + [float(r["返回O01时刻_s"]) for r in relays]), "联合完成")
    assert metrics["transport_sorties"] == 22 and metrics["relay_sorties"] == 4
    assert metrics["boxes"] == 80 and metrics["hard_excess_s"] == 0
    assert metrics["weighted_tardiness"] == 0 and metrics["communication_gap_s"] == 0
    return full


def verify_q4(q3):
    areas = {stop["area"] for trip in q3["transport"] for stop in trip["stops"]}
    assert len(areas) == 15
    relay_use = defaultdict(set)
    for row in q3["communications"]:
        if row["relay_id"]:
            relay_use[row["trip"]].add(row["relay_id"])
    relays = {r["id"]: r for r in q3["relays"]}
    output = []
    for k in (2, 3):
        for name in ("缺口优先", "兼顾均衡"):
            plan = json.loads((Q4 / f"K{k}_{name}.json").read_text(encoding="utf-8"))
            groups = [set(a) for a in plan["任务组"]]
            assert len(groups) == k and all(groups) and set.union(*groups) == areas
            assert sum(len(g) for g in groups) == 15
            facts, work = [], []
            for group in groups:
                trips = [t for t in q3["transport"] if t["stops"][0]["area"] in group]
                assert all({s["area"] for s in t["stops"]} <= group for t in trips)
                used = set().union(*(relay_use[t["id"]] for t in trips)) if trips else set()
                work.append(sum(t["return_s"] - t["start_s"] for t in trips))
                row = {}
                for model in "ABC":
                    selected = [t for t in trips if t["model"] == model]
                    row["U_" + model] = peak([(t["start_s"], t["return_s"]) for t in selected])
                    row["B_" + model] = peak([(t["start_s"], t["charge_end_s"]) for t in selected])
                selected_relays = [relays[rid] for rid in used]
                row["R"] = peak([(r["depart_s"], r["relay_free_s"]) for r in selected_relays])
                row["RB"] = peak([(r["depart_s"], r["energy_free_s"]) for r in selected_relays])
                facts.append(row)
            assert facts == plan["资源需求"]
            total = {category: sum(row[category] for row in facts) for category in STOCK}
            short = {category: max(0, total[category] - stock) for category, stock in STOCK.items()}
            assert total == plan["资源总需求"] and short == plan["资源缺口"]
            assert sum(short.values()) == plan["缺口总数"]
            for a, b in zip(work, plan["各组运输作业量_s"]):
                close(a, b, (k, name, "作业量"))
            mean = sum(work) / k
            cv = (sum((v - mean) ** 2 for v in work) / k) ** .5 / mean
            close(cv, plan["作业量CV"], (k, name, "CV"))
            output.append({"K": k, "方案": name, "缺口": sum(short.values()), "CV": cv})
    return output


def verify_template():
    book = load_workbook(Q4 / "结果提交_问题三四主方案.xlsx", read_only=True, data_only=True)
    prior = load_workbook(ROOT / "results" / "问题二_参考口径" / "结果提交_主方案.xlsx",
                          read_only=True, data_only=True)
    assert book.sheetnames == ["Q1_单点组批", "Q2_运输架次", "Q2_逐箱交付",
                               "Q3_中继架次", "Q3_通信保障", "Q4_分区配置"]
    assert list(book["Q1_单点组批"].values) == list(prior["Q1_单点组批"].values)

    def same(cell, raw):
        if cell is None:
            return raw == ""
        if isinstance(cell, (int, float)):
            try:
                return math.isclose(float(cell), float(raw), rel_tol=1e-9, abs_tol=1e-8)
            except ValueError:
                return False
        return str(cell) == raw

    for sheet, csv_file, expected, count in (("Q2_运输架次", "主方案_运输架次.csv", 22, 8),
                                             ("Q2_逐箱交付", "主方案_逐箱交付.csv", 80, 4),
                                             ("Q3_中继架次", "主方案_中继架次.csv", 4, 11),
                                             ("Q3_通信保障", "主方案_通信保障.csv", None, 6)):
        rows = read(Q3 / csv_file)
        assert book[sheet].max_row == len(rows) + 1
        if expected is not None:
            assert len(rows) == expected
        columns = list(rows[0])[:count]
        cells = list(book[sheet].values)
        for index, (csv_row, xlsx_row) in enumerate(zip(rows, cells[1:]), 2):
            assert all(same(cell, csv_row[key]) for key, cell in zip(columns, xlsx_row[:count])), (sheet, index)
            assert all(cell is None for cell in xlsx_row[count:]), (sheet, index, "多余单元格")
    q4_rows = read(Q4 / "问题四主方案逐组资源.csv")
    assert len(q4_rows) == 5
    columns = list(q4_rows[0])[:11]
    cells = list(book["Q4_分区配置"].values)
    for index, (csv_row, xlsx_row) in enumerate(zip(q4_rows, cells[1:]), 2):
        assert all(same(cell, csv_row[key]) for key, cell in zip(columns, xlsx_row[:11])), ("Q4", index)
        assert all(cell is None for cell in xlsx_row[11:]), ("Q4", index, "多余单元格")
    book.close()
    prior.close()


def main():
    source = inputs()
    q3 = verify_q3("主方案", source)
    q4 = verify_q4(q3)
    verify_template()
    report = {"状态": "PASS", "问题三主方案指标": q3["metrics"], "问题四复算": q4,
              "审计范围": "逐箱、运输航段/能量/资源、通信表连续覆盖与中继窗口、两阶段充电、分区峰值/库存/工作量；DEM 链路使用求解器保守区间证书另审"}
    (Q4 / "独立审计报告.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False))


if __name__ == "__main__":
    main()
