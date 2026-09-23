"""D题第一问：DEM 单点往返载荷与类型压缩组批整数规划。

从本文件所在目录运行：python 问题1_求解.py --smoke S001
完整运行：python 问题1_求解.py
"""
from __future__ import annotations

import argparse
import csv
import itertools
import json
import math
from collections import Counter, defaultdict, deque
from dataclasses import dataclass
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.io import loadmat
from scipy.optimize import Bounds, LinearConstraint, milp

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "数据" / "无人机应急物资运输基础数据"
DEM_PATH = next((ROOT / "数据").rglob("*DEM.mat"))
G = 9.80665


class InfeasibleArea(RuntimeError):
    """给定安全余量下某服务区货箱无法全部组批。"""


@dataclass(frozen=True)
class Drone:
    name: str
    mass0: float
    payload: float
    volume: float
    speed: float
    range0: float
    range_full: float
    energy: float
    reserve: float
    prep: float
    load: float
    handoff: float
    extra: float
    climb_speed: float
    descent_speed: float
    eta: float


@dataclass(frozen=True)
class Geometry:
    area: str
    distance: float
    cruise: float
    climb_out: float
    descent_out: float
    climb_back: float
    descent_back: float
    cells: int


def rows(path: Path, sheet: int = 0):
    book = load_workbook(path, read_only=True, data_only=True)
    return list(book.worksheets[sheet].values)


def inputs():
    drone_rows = rows(DATA / "运输无人机数据.xlsx")
    drones = {}
    for row in drone_rows[2:5]:
        x = row[:18]
        drones[x[0]] = Drone(x[0], *map(float, x[2:10]), *map(float, x[10:17]))
    assert set(drones) == {"A", "B", "C"}
    node_rows = rows(DATA / "调度中心与服务区.xlsx")
    origin = node_rows[2]
    assert origin[0] == "O01"
    sites = {r[0]: (float(r[2]), float(r[3]), float(r[4])) for r in node_rows[6:21]}
    assert len(sites) == 15
    box_rows = rows(DATA / "物资需求与配送时限.xlsx", 1)
    assert len(box_rows) == 81 and box_rows[0][0] == "货箱编号" and box_rows[-1][0] == "S015-FOD-01"
    boxes = defaultdict(list)
    seen = set()
    for r in box_rows[1:]:
        assert r[0] not in seen and r[1] in sites and r[3] > 0 and r[4] > 0
        seen.add(r[0])
        boxes[r[1]].append({"id": r[0], "type": r[2], "mass": float(r[3]), "volume": float(r[4])})
    assert len(seen) == 80
    demand_rows = rows(DATA / "物资需求与配送时限.xlsx", 0)
    summary = {(r[0], r[1]): int(r[2]) for r in demand_rows[1:] if r[0]}
    assert Counter((s, b["type"]) for s, bs in boxes.items() for b in bs) == summary
    return drones, (float(origin[2]), float(origin[3]), float(origin[4])), sites, boxes


def load_dem():
    mat = loadmat(DEM_PATH)
    dem = mat["dem"]
    lat = mat["latitude"][:, 0]
    lon = mat["longitude"][0]
    nodata = float(mat["nodata"][0, 0])
    assert dem.shape == (1309, 1486) and int(mat["epsg_code"][0, 0]) == 4326
    return dem, float(lat[0]), float(lon[0]), float(lon[1] - lon[0]), float(lat[0] - lat[1]), nodata


def line_cells(lon0, lat0, lon1, lat1, dem_meta):
    dem, lat_top, lon_left, dx, dy, nodata = dem_meta
    x0, y0 = (lon0 - lon_left) / dx, (lat_top - lat0) / dy
    x1, y1 = (lon1 - lon_left) / dx, (lat_top - lat1) / dy
    breaks = [0.0, 1.0]
    for a, b in ((x0, x1), (y0, y1)):
        if abs(b - a) < 1e-12:
            continue
        lo, hi = sorted((a, b))
        for boundary in np.arange(math.floor(lo - 0.5) + 1, math.ceil(hi - 0.5) + 1) + 0.5:
            t = (boundary - a) / (b - a)
            if 0 < t < 1:
                breaks.append(float(t))
    breaks.sort()
    values = []
    for left, right in zip(breaks[:-1], breaks[1:]):
        t = (left + right) / 2
        row = math.floor(y0 + t * (y1 - y0) + 0.5)
        col = math.floor(x0 + t * (x1 - x0) + 0.5)
        assert 0 <= row < dem.shape[0] and 0 <= col < dem.shape[1], "航段离开DEM覆盖"
        z = float(dem[row, col])
        assert z != nodata and np.isfinite(z), "航段经过无效DEM像元"
        values.append(z)
    return values


def haversine(lon0, lat0, lon1, lat1):
    p0, p1 = math.radians(lat0), math.radians(lat1)
    dp, dl = p1 - p0, math.radians(lon1 - lon0)
    a = math.sin(dp / 2) ** 2 + math.cos(p0) * math.cos(p1) * math.sin(dl / 2) ** 2
    return 2 * 6371008.8 * math.asin(math.sqrt(a))


def geometry(area, origin, site, dem_meta):
    lo, la, zo = origin
    ls, lt, zs = site
    zmax = max(line_cells(lo, la, ls, lt, dem_meta))
    cruise = zmax + 50
    assert cruise >= max(zo, zs + 30), (area, cruise, zo, zs)
    return Geometry(area, haversine(lo, la, ls, lt), cruise,
                    cruise - zo, cruise - (zs + 30),
                    cruise - (zs + 30), cruise - zo,
                    len(line_cells(lo, la, ls, lt, dem_meta)))


def equivalent_range(drone, q):
    assert 0 <= q <= drone.payload + 1e-8
    return drone.range0 - (drone.range0 - drone.range_full) * (q / drone.payload) ** 1.5


def energy(drone, geo, q):
    horizontal_out = drone.energy * geo.distance / equivalent_range(drone, q)
    horizontal_back = drone.energy * geo.distance / drone.range0
    climb_out = (drone.mass0 + q) * G * geo.climb_out / (3.6e6 * drone.eta)
    climb_back = drone.mass0 * G * geo.climb_back / (3.6e6 * drone.eta)
    return horizontal_out + horizontal_back + climb_out + climb_back


def flight_time(drone, geo):
    return ((geo.climb_out + geo.climb_back) / drone.climb_speed
            + 2 * geo.distance / drone.speed
            + (geo.descent_out + geo.descent_back) / drone.descent_speed)


def trip_time(drone, geo, n):
    return flight_time(drone, geo) + drone.prep + n * drone.load + drone.handoff + n * drone.extra


def safe_payload(drone, geo, reserve):
    limit = (1 - reserve) * drone.energy
    if energy(drone, geo, 0) > limit + 1e-10:
        return None
    if energy(drone, geo, drone.payload) <= limit:
        return drone.payload
    lo, hi = 0.0, drone.payload
    for _ in range(55):
        mid = (lo + hi) / 2
        if energy(drone, geo, mid) <= limit:
            lo = mid
        else:
            hi = mid
    return lo


def patterns(area_boxes, drones, geo, reserve):
    by_type = defaultdict(list)
    for b in area_boxes:
        by_type[b["type"]].append(b)
    types = sorted(by_type)
    attrs = []
    for k in types:
        bs = by_type[k]
        assert all((b["mass"], b["volume"]) == (bs[0]["mass"], bs[0]["volume"]) for b in bs)
        attrs.append((len(bs), bs[0]["mass"], bs[0]["volume"]))
    pats = []
    for counts in itertools.product(*(range(a[0] + 1) for a in attrs)):
        n = sum(counts)
        if n == 0:
            continue
        mass = sum(c * a[1] for c, a in zip(counts, attrs))
        volume = sum(c * a[2] for c, a in zip(counts, attrs))
        for name, drone in drones.items():
            if mass > drone.payload + 1e-9 or volume > drone.volume + 1e-9:
                continue
            e = energy(drone, geo, mass)
            if e <= (1 - reserve) * drone.energy + 1e-9:
                pats.append({"counts": counts, "model": name, "mass": mass,
                             "volume": volume, "energy": e,
                             "time": trip_time(drone, geo, n), "n": n})
    return types, attrs, pats


def solve_stage(pats, demand, objective, n_fixed=None, e_cap=None):
    from scipy.sparse import csc_array
    columns = len(pats)
    A = np.array([p["counts"] for p in pats], dtype=float).T
    lower = np.asarray(demand, dtype=float)
    upper = lower.copy()
    if n_fixed is not None:
        A = np.vstack([A, np.ones(columns)])
        lower = np.r_[lower, n_fixed]
        upper = np.r_[upper, n_fixed]
    if e_cap is not None:
        A = np.vstack([A, [p["energy"] for p in pats]])
        lower = np.r_[lower, -np.inf]
        upper = np.r_[upper, e_cap]
    c = np.array([p[objective] if objective != "trips" else 1.0 for p in pats])
    result = milp(c, integrality=np.ones(columns), bounds=Bounds(0, np.inf),
                  constraints=LinearConstraint(csc_array(A), lower, upper),
                  options={"time_limit": 120, "mip_rel_gap": 1e-9})
    if result.status == 2:
        raise InfeasibleArea("货箱覆盖整数规划不可行")
    if result.status != 0 or result.x is None:
        raise RuntimeError(f"整数规划未证最优: {result.message}")
    x = np.rint(result.x).astype(int)
    assert np.allclose(A[:len(demand)] @ x, demand, atol=1e-7)
    return x, float(c @ x)


def optimize_area(area_boxes, drones, geo, reserve, alternative=False):
    types, attrs, pats = patterns(area_boxes, drones, geo, reserve)
    if not pats:
        raise InfeasibleArea(f"{geo.area}: 没有可行组批")
    demand = [a[0] for a in attrs]
    x1, nopt = solve_stage(pats, demand, "trips")
    x2, eopt = solve_stage(pats, demand, "energy", n_fixed=round(nopt))
    x3, _ = solve_stage(pats, demand, "time", n_fixed=round(nopt), e_cap=eopt + 1e-8)
    selected = [p for p, x in zip(pats, x3) for _ in range(x)]
    assert len(selected) == round(nopt)
    alt = None
    if alternative:
        xa, _ = solve_stage(pats, demand, "energy")
        alt = {"trips": int(sum(xa)), "energy": float(sum(x * p["energy"] for x, p in zip(xa, pats))),
               "time": float(sum(x * p["time"] for x, p in zip(xa, pats)))}
    return types, selected, alt


def assign_boxes(types, selected, area_boxes):
    pools = {t: deque(sorted((b["id"] for b in area_boxes if b["type"] == t))) for t in types}
    out = []
    for p in selected:
        ids = []
        for t, n in zip(types, p["counts"]):
            for _ in range(n):
                ids.append(pools[t].popleft())
        out.append((p, ids))
    assert all(not x for x in pools.values())
    return out


def write_csv(path, columns, records):
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=columns)
        w.writeheader()
        w.writerows(records)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--smoke", choices=[f"S{i:03d}" for i in range(1, 16)])
    args = parser.parse_args()
    drones, origin, sites, boxes = inputs()
    dem = load_dem()
    if args.smoke:
        area = args.smoke
        geo = geometry(area, origin, sites[area], dem)
        caps = {g: safe_payload(d, geo, d.reserve / 100) for g, d in drones.items()}
        types, chosen, alt = optimize_area(boxes[area], drones, geo, 0.2, True)
        assigned = assign_boxes(types, chosen, boxes[area])
        assert sum(len(ids) for _, ids in assigned) == len(boxes[area])
        print(json.dumps({"area": area, "distance_m": geo.distance, "cruise_m": geo.cruise,
                          "safe_payload_kg": caps, "trips": len(chosen),
                          "energy_kwh": sum(p["energy"] for p in chosen),
                          "operation_s": sum(p["time"] for p in chosen),
                          "alternative_energy_first": alt,
                          "boxes_assigned": sum(len(ids) for _, ids in assigned)}, ensure_ascii=False, indent=2))
        return
    geometries = {s: geometry(s, origin, site, dem) for s, site in sites.items()}
    caps_rows = []
    for s, geo in geometries.items():
        for g, d in drones.items():
            cap = safe_payload(d, geo, d.reserve / 100)
            caps_rows.append({"服务区编号": s, "机型编号": g,
                              "水平距离（m）": geo.distance, "DEM最高高程（m）": geo.cruise - 50,
                              "计划巡航海拔（m）": geo.cruise, "经过DEM像元数": geo.cells,
                              "最大安全载荷（kg）": cap,
                              "满载能耗（kWh）": energy(d, geo, d.payload),
                              "空载能耗（kWh）": energy(d, geo, 0)})
    trips = []
    alternatives = []
    for s, geo in geometries.items():
        types, selected, alt = optimize_area(boxes[s], drones, geo, 0.2, True)
        for p, ids in assign_boxes(types, selected, boxes[s]):
            d = drones[p["model"]]
            trips.append({"架次编号": f"Q1-{len(trips)+1:03d}", "服务区编号": s,
                          "机型编号": d.name, "货箱编号列表": ",".join(ids),
                          "总质量（kg）": p["mass"], "总体积（m³）": p["volume"],
                          "往返时间（s）": flight_time(d, geo), "作业时间（s）": p["time"],
                          "架次能耗（kWh）": p["energy"],
                          "返航SOC（%）": 100 * (1 - p["energy"] / d.energy)})
        alternatives.append({"服务区编号": s, "方案": "架次→能耗→时间", "架次数": len(selected),
                             "总能耗（kWh）": sum(p["energy"] for p in selected),
                             "累计作业时间（s）": sum(p["time"] for p in selected)})
        alternatives.append({"服务区编号": s, "方案": "能耗优先", "架次数": alt["trips"],
                             "总能耗（kWh）": alt["energy"], "累计作业时间（s）": alt["time"]})
    assert len({bid for trip in trips for bid in trip["货箱编号列表"].split(",")}) == 80
    assert sum(len(t["货箱编号列表"].split(",")) for t in trips) == 80
    for trip in trips:
        d = drones[trip["机型编号"]]
        assert trip["总质量（kg）"] <= d.payload + 1e-8
        assert trip["总体积（m³）"] <= d.volume + 1e-8
        assert trip["返航SOC（%）"] >= d.reserve - 1e-7
    sensitivity = []
    sensitivity_trips = []
    for reserve in (0.10, 0.20, 0.30, 0.40, 0.50):
        total_n = total_e = total_t = 0.0
        scenario_trips = []
        for s, geo in geometries.items():
            for g, d in drones.items():
                sensitivity.append({"返航余量（%）": reserve * 100, "服务区编号": s, "机型编号": g,
                                    "最大安全载荷（kg）": safe_payload(d, geo, reserve)})
            try:
                types, selected, _ = optimize_area(boxes[s], drones, geo, reserve)
                for row in sensitivity[-3:]:
                    row["该服务区可完成"] = True
                total_n += len(selected)
                total_e += sum(p["energy"] for p in selected)
                total_t += sum(p["time"] for p in selected)
                for p, ids in assign_boxes(types, selected, boxes[s]):
                    scenario_trips.append({"返航余量（%）": reserve * 100, "服务区编号": s,
                                           "机型编号": p["model"], "货箱编号列表": ",".join(ids),
                                           "总质量（kg）": p["mass"], "总体积（m³）": p["volume"],
                                           "架次能耗（kWh）": p["energy"], "作业时间（s）": p["time"]})
            except InfeasibleArea:
                for row in sensitivity[-3:]:
                    row["该服务区可完成"] = False
                total_n = total_e = total_t = float("nan")
        for row in sensitivity[-45:]:
            row["全任务架次"] = total_n
            row["全任务能耗（kWh）"] = total_e
            row["全任务累计时间（s）"] = total_t
        if np.isfinite(total_n):
            assert sum(len(r["货箱编号列表"].split(",")) for r in scenario_trips) == 80
            sensitivity_trips.extend(scenario_trips)
    out = ROOT / "results" / "历史结果_旧口径"
    out.mkdir(parents=True, exist_ok=True)
    write_csv(out / "最大安全载荷.csv", list(caps_rows[0]), caps_rows)
    write_csv(out / "单点组批.csv", list(trips[0]), trips)
    write_csv(out / "指标权衡.csv", list(alternatives[0]), alternatives)
    write_csv(out / "返航余量敏感性.csv", list(sensitivity[0]), sensitivity)
    write_csv(out / "返航余量组批.csv", list(sensitivity_trips[0]), sensitivity_trips)
    template = load_workbook(ROOT / "结果提交模板.xlsx")
    sheet = template["Q1_单点组批"]
    template_columns = [sheet.cell(1, j).value for j in range(1, 10)]
    for i, trip in enumerate(trips, 2):
        for j, key in enumerate(template_columns, 1):
            sheet.cell(i, j, trip[key])
    template.save(out / "第一问_模板填写结果.xlsx")
    print(json.dumps({"trips": len(trips), "boxes": 80,
                      "energy_kwh": sum(t["架次能耗（kWh）"] for t in trips),
                      "operation_s": sum(t["作业时间（s）"] for t in trips),
                      "min_return_soc_pct": min(t["返航SOC（%）"] for t in trips)},
                     ensure_ascii=False, indent=2))
    from 问题1_绘图 import draw
    draw()


if __name__ == "__main__":
    main()
