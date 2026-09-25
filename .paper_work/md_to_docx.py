import json
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


ROOT = Path(__file__).resolve().parents[1]
META = json.loads((ROOT / '.paper_work' / 'paper_meta.json').read_text(encoding='utf-8'))
PARTS = ['paper_01.md', 'paper_02.md', 'paper_03.md', 'paper_04.md', 'paper_05.md', 'paper_06.md', 'paper_07.md']
MD = '\n'.join((ROOT / '.paper_work' / p).read_text(encoding='utf-8') for p in PARTS)
TEMPLATE = ROOT / '.paper_work' / '华为杯论文模板.docx'
OUT = ROOT / '山区洪涝灾害下无人机运输与通信协同优化_论文.docx'

CN_FONT = '宋体'
HEAD_FONT = '黑体'
MATH_FONT = 'Cambria Math'
MONO_FONT = 'Consolas'


def set_run(run, size=10.5, name=CN_FONT, bold=False, color=None, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = name
    r = run._element.rPr.rFonts
    r.set(qn('w:eastAsia'), name)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def replace_text(par, text, size=None, bold=None, name=None):
    if par.runs:
        par.runs[0].text = text
        for run in par.runs[1:]:
            run.text = ''
        target = par.runs[0]
    else:
        target = par.add_run(text)
    if size is not None or bold is not None or name is not None:
        set_run(target, size=size if size is not None else 10.5,
                bold=bold if bold is not None else False,
                name=name if name is not None else CN_FONT)


def add_body_paragraph(doc, text, indent=True, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                       size=10.5, name=CN_FONT, bold=False, space_after=0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.line_spacing = 1.25
    pf.space_after = Pt(space_after)
    pf.first_line_indent = Pt(21) if indent else Pt(0)
    p.alignment = align
    # simple bold parsing
    pos = 0
    for m in re.finditer(r'\*\*(.+?)\*\*', text):
        if m.start() > pos:
            r = p.add_run(text[pos:m.start()]); set_run(r, size=size, name=name, bold=bold)
        r = p.add_run(m.group(1)); set_run(r, size=size, name=name, bold=True)
        pos = m.end()
    if pos < len(text):
        r = p.add_run(text[pos:]); set_run(r, size=size, name=name, bold=bold)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(12 if level == 1 else 8)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    if level == 1:
        size = 15
    elif level == 2:
        size = 12.5
    else:
        size = 11
    r = p.add_run(text.strip('# ').strip())
    set_run(r, size=size, name=HEAD_FONT, bold=True)
    return p


def add_formula(doc, lines):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    text = ' '.join(x.strip() for x in lines).strip('$').strip()
    r = p.add_run(text)
    set_run(r, size=10.5, name=MATH_FONT)
    return p


def add_image(doc, alt, path):
    img = ROOT / path
    if not img.exists():
        return
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    try:
        p.add_run().add_picture(str(img), width=Cm(14.2))
    except Exception:
        return
    cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(6)
    r = cap.add_run(alt)
    set_run(r, size=9, name=CN_FONT)


def add_table(doc, rows):
    if not rows:
        return
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            text = row[j] if j < len(row) else ''
            cell.text = ''
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(text)
            set_run(r, size=8.5, name=CN_FONT, bold=(i == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def parse_markdown(doc, text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        # code block
        if stripped.startswith('```'):
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i].rstrip()); i += 1
            i += 1
            for cl in code:
                p = doc.add_paragraph()
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.space_after = Pt(0)
                r = p.add_run(cl if cl else ' ')
                set_run(r, size=9, name=MONO_FONT)
            continue
        # image
        m = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
        if m:
            add_image(doc, m.group(1), m.group(2)); i += 1; continue
        # formula block
        if stripped.startswith('$$'):
            block = [stripped]
            if stripped.count('$$') < 2:
                i += 1
                while i < len(lines):
                    block.append(lines[i].strip())
                    if lines[i].strip().endswith('$$'):
                        break
                    i += 1
            add_formula(doc, block); i += 1; continue
        # table
        if stripped.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                raw = lines[i].strip().strip('|')
                cells = [c.strip() for c in raw.split('|')]
                if all(set(c) <= set('-: ') for c in cells):
                    i += 1; continue
                rows.append(cells); i += 1
            add_table(doc, rows)
            continue
        # headings
        if stripped.startswith('### '):
            add_heading(doc, stripped[4:], 3); i += 1; continue
        if stripped.startswith('## '):
            add_heading(doc, stripped[3:], 2); i += 1; continue
        if stripped.startswith('# '):
            add_heading(doc, stripped[2:], 1); i += 1; continue
        # list bullet
        if stripped.startswith('- '):
            p = add_body_paragraph(doc, '· ' + stripped[2:], indent=False)
            p.paragraph_format.left_indent = Pt(21)
            i += 1; continue
        # normal paragraph
        add_body_paragraph(doc, stripped, indent=True)
        i += 1


def main():
    doc = Document(str(TEMPLATE))
    ps = doc.paragraphs
    # cover title and abstract
    for p in ps:
        if '题 目' in p.text:
            replace_text(p, '题 目：' + META['title'])
        if p.text.strip().startswith('摘 要'):
            replace_text(p, '摘 要：', size=14, bold=True, name='黑体')
        if p.text.strip().startswith('关键词'):
            replace_text(p, '关键词：' + META['keywords'], size=10.5, bold=False, name=CN_FONT)
    # fill abstract into the paragraph after 摘 要
    idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith('摘 要'):
            idx = i; break
    if idx is not None and idx + 1 < len(doc.paragraphs):
        ap = doc.paragraphs[idx + 1]
        replace_text(ap, META['abstract'], size=10.5, bold=False, name=CN_FONT)
        ap.paragraph_format.first_line_indent = Pt(0)
        ap.paragraph_format.line_spacing = 1.25
    # remove remaining blank template paragraphs after abstract area (keep first 24)
    for p in list(doc.paragraphs)[24:]:
        p._element.getparent().remove(p._element)
    # page break
    pb = doc.add_paragraph(); run = pb.add_run(); run.add_break(WD_BREAK.PAGE)
    parse_markdown(doc, MD)
    doc.save(str(OUT))
    print('saved', OUT)


if __name__ == '__main__':
    main()
