"""t5 论文参考文献口径修正：[9]/[10] 按题面 题目原文.txt:98-99 校正，原两条顺延为 [11]/[12]。

题面出处：proj-d-uav-flood/题目原文.txt:98-99
  [9] DJI Enterprise. Matrice 350 RTK Specifications[EB/OL]. https://enterprise.dji.com/matrice-350-rtk/specs.
  [10] Doodle Labs. Sense: Interference Avoidance[EB/OL]. https://doodlelabs.com/news/sense-interference-avoidance-release/
原 [9] Ropke（式(2-2)/(2-3) 之外的大邻域搜索方法来源）与原 [10]（D题附件与数据文件）不删除，顺延为 [11]/[12]。
正文未出现 [9]/[10] 引用标记（已用正则核对 .paper_work/paper_0*.md 与 docx 正文），故无需改标参数取值来源。
docx 由 .paper_work/md_to_docx.py 无法字节级复现（现行脚本会引入 '# ' 与表格 XML 残留），因此对 docx 做等效的最小 XML/段落级同步。
"""
from __future__ import annotations

import copy
import hashlib
import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[2]
MD = ROOT / ".paper_work" / "paper_07.md"
DOCX = ROOT / "山区洪涝灾害下无人机运输与通信协同优化_论文.docx"

OLD9 = "[9] Ropke S, Pisinger D. An Adaptive Large Neighborhood Search Heuristic for the Pickup and Delivery Problem with Time Windows[J]. Transportation Science, 2006, 40(4):455-472."
OLD10 = "[10] 华为杯第二十三届中国研究生数学建模竞赛D题附件与数据文件."
NEW9 = "[9] DJI Enterprise. Matrice 350 RTK Specifications[EB/OL]. https://enterprise.dji.com/matrice-350-rtk/specs."
NEW10 = "[10] Doodle Labs. Sense: Interference Avoidance[EB/OL]. https://doodlelabs.com/news/sense-interference-avoidance-release/."
NEW11 = "[11] Ropke S, Pisinger D. An Adaptive Large Neighborhood Search Heuristic for the Pickup and Delivery Problem with Time Windows[J]. Transportation Science, 2006, 40(4):455-472."
NEW12 = "[12] 华为杯第二十三届中国研究生数学建模竞赛D题附件与数据文件."


def patch_md() -> None:
    lines = MD.read_text(encoding="utf-8").splitlines()
    i9 = next(i for i, line in enumerate(lines) if line.strip() == OLD9)
    i10 = next(i for i, line in enumerate(lines) if line.strip() == OLD10)
    assert i10 == i9 + 2, (i9, i10)
    lines[i9], lines[i10] = NEW9, NEW10
    lines[i10 + 1:i10 + 1] = ["", NEW11, "", NEW12]
    MD.write_text("\n".join(lines) + "\n", encoding="utf-8")
    print(f"paper_07.md: {i9 + 1}/{i10 + 1} 行更新，新增 [11]/[12]")
    print(f"  SHA-256={hashlib.sha256(MD.read_bytes()).hexdigest()}")


def set_text(par: Paragraph, text: str) -> None:
    if par.runs:
        par.runs[0].text = text
        for run in par.runs[1:]:
            run.text = ""
    else:
        par.add_run(text)


def patch_docx() -> None:
    doc = Document(str(DOCX))
    refs = {}
    for par in doc.paragraphs:
        match = re.match(r"^\[(\d+)\]", par.text.strip())
        if match:
            refs[int(match.group(1))] = par
    assert 9 in refs and 10 in refs, sorted(refs)
    before = {"[9]": refs[9].text.strip(), "[10]": refs[10].text.strip()}
    set_text(refs[9], NEW9)
    anchor = refs[10]
    set_text(anchor, NEW10)
    new_elements = []
    for text in (NEW11, NEW12):
        element = copy.deepcopy(anchor._p)
        anchor._p.addnext(element)
        new_elements.append(Paragraph(element, anchor._parent))
        anchor = new_elements[-1]
        set_text(anchor, text)
    doc.save(str(DOCX))
    print("docx 原值:", before)
    print(f"docx 新值: {{'[9]': {NEW9[:40]}…, '[10]': {NEW10[:40]}…, '[11]/[12] 已补}}")
    print(f"  SHA-256={hashlib.sha256(DOCX.read_bytes()).hexdigest()}")


def verify() -> None:
    md_refs = [line.strip() for line in MD.read_text(encoding="utf-8").splitlines()
               if re.match(r"^\[\d+\] ", line.strip())]
    doc = Document(str(DOCX))
    docx_refs = [par.text.strip() for par in doc.paragraphs if re.match(r"^\[\d+\] ", par.text.strip())]
    print("md 文献条数", len(md_refs), "docx 文献条数", len(docx_refs))
    for i, (a, b) in enumerate(zip(md_refs, docx_refs), 1):
        assert a == b, (i, a, b)
    print("md 与 docx 参考文献逐条一致：True")
    print("docx 尾三条:", docx_refs[-3:])


if __name__ == "__main__":
    patch_md()
    patch_docx()
    verify()
