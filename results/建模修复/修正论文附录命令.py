"""t5 修正论文附录A 复现命令：问题三产物写入入口是 问题三_求解.py（默认 --resolution 0.5，内部调用 问题三_结果输出.save）；问题三_联合调度.py main() 只打印诊断指标、不写产物。"""
from __future__ import annotations

import hashlib
import re
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[2]
MD = ROOT / ".paper_work" / "paper_07.md"
DOCX = ROOT / "山区洪涝灾害下无人机运输与通信协同优化_论文.docx"
OLD = "python 问题三_联合调度.py --resolution 0.5"
NEW = "python 问题三_求解.py --resolution 0.5"


def main() -> None:
    text = MD.read_text(encoding="utf-8")
    assert text.count(OLD) == 1, text.count(OLD)
    line = next(i for i, value in enumerate(text.splitlines(), 1) if OLD in value)
    MD.write_text(text.replace(OLD, NEW), encoding="utf-8")
    print(f"[paper_07.md:{line}] {OLD}  ->  {NEW}")
    print(f"  SHA-256={hashlib.sha256(MD.read_bytes()).hexdigest()}")

    doc = Document(str(DOCX))
    hit = [p for p in doc.paragraphs if p.text.strip() == OLD]
    assert len(hit) == 1, len(hit)
    par = hit[0]
    if par.runs:
        par.runs[0].text = NEW
        for run in par.runs[1:]:
            run.text = ""
    else:
        par.add_run(NEW)
    doc.save(str(DOCX))
    print(f"[docx] {OLD}  ->  {NEW}")
    print(f"  SHA-256={hashlib.sha256(DOCX.read_bytes()).hexdigest()}")

    check = [p.text.strip() for p in Document(str(DOCX)).paragraphs if re.match(r"^python ", p.text.strip())]
    print("docx 复现命令块:", check)


if __name__ == "__main__":
    main()
